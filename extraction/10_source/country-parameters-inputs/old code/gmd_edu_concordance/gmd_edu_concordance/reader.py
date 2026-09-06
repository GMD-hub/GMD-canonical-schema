"""
gmd_edu_concordance.reader
==========================
UIS "ISCED 2011 Mapping" workbook -> Instruction.

What the workbook actually holds
--------------------------------
One banner row naming the country, one school-year reference, and a table of
national programmes.  Per programme: the name in both languages, the minimum
entrance requirement, the qualification awarded, the **theoretical entrance
age**, the **theoretical duration in years**, the ISCED 2011 level and its
code, orientation, completion, access, and the three-digit ISCED-P and ISCED-A
codes.

What it does not hold
---------------------
Grades, and any structure other than the one in force in the reference year.
Both are worked out here and marked inferred; see ``derive.py``.
"""
from __future__ import annotations

import os
import re
from typing import Any, Dict, List, Optional, Tuple

import yaml
from openpyxl import load_workbook

from . import countries as CO
from . import parse as P
from .schema import (
    ISCED_LEVELS, Instruction, SourceFile, EduRow, Unmapped, inferred,
)

DATA = os.path.join(os.path.dirname(__file__), "data")
DEFAULT_PROFILE = "isced_uis_2011_mapping.yaml"


def load_profile(path: Optional[str] = None) -> dict:
    with open(path or os.path.join(DATA, DEFAULT_PROFILE), encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def load_profiles(path: Optional[str] = None) -> List[dict]:
    """Every binding profile, most specific first.

    More than one workbook family carries an ISCED mapping. The UIS national
    mapping sheet is one; the UOE questionnaire several European countries
    return instead is another. Which one a file is is decided by trying to bind
    it, not by its name -- so a new family is a YAML file, not a second reader.
    """
    if path:
        return [load_profile(path)]
    out = []
    for f in sorted(os.listdir(DATA)):
        if not f.endswith(".yaml") or f == "iso3166.yaml":
            continue
        with open(os.path.join(DATA, f), encoding="utf-8") as fh:
            d = yaml.safe_load(fh)
        if isinstance(d, dict) and d.get("profile") and d.get("columns"):
            out.append(d)
    out.sort(key=lambda d: d.get("priority", 100))
    return out


def _txt(ws, r, c) -> str:
    v = ws.cell(r, c).value
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


class Binding:
    """What bound where.  ``inspect`` prints this and stops."""

    def __init__(self, profile_name: str):
        self.profile = profile_name
        self.sheet = ""
        self.header_row: Optional[int] = None
        self.columns: Dict[str, Optional[int]] = {}
        self.headers: Dict[str, str] = {}
        self.identity: Dict[str, Any] = {}
        self.unbound: List[str] = []
        self.candidates: Dict[str, List[Tuple[int, str]]] = {}
        self.problems: List[str] = []
        self.tried: List[str] = []
        self.candidates_considered: List[dict] = []
        self.score: Optional[dict] = None
        self.rows_read = 0

    def as_dict(self):
        return {"profile": self.profile, "sheet": self.sheet,
                "header_row": self.header_row, "identity": self.identity,
                "columns": self.columns, "headers": self.headers,
                "rows_read": self.rows_read, "unbound": self.unbound,
                "profiles_tried": self.tried,
                "sheets_considered": self.candidates_considered,
                "chosen": self.score,
                "candidates": {k: [{"col": c, "header": h} for c, h in v]
                               for k, v in self.candidates.items()},
                "problems": self.problems}


# --------------------------------------------------------------------------
def _looks_like_prose(ws, header_row: int) -> bool:
    """A sheet of instructions has long sentences where a table has labels."""
    long_cells = 0
    for r in range(1, min(header_row, 12) + 1):
        for c in range(1, 4):
            v = _txt(ws, r, c)
            if len(v) > 160:
                long_cells += 1
    return long_cells >= 3


def score_sheet(ws, prof, name: str) -> Optional[dict]:
    """How well does this sheet look like this profile's table?

    A workbook often holds several sheets and only one of them is the mapping:
    an instructions page, a dropdown-list page, a per-level annex, and the
    table itself. Taking the first sheet whose header row happens to match
    enough strings is how the wrong one gets read, and the failure is silent --
    it produces a short, plausible instruction rather than an error.

    So every sheet is scored against every profile and the best pair wins.
    What counts, in order:

      * do the profile's **required** columns bind at all -- without them the
        sheet cannot be the table, whatever else matches;
      * how many **rows actually parse** to a valid ISCED level -- the
        strongest signal, because an instructions page has none;
      * how many columns bound, as a tie-break.
    """
    hdr = _find_header(ws, prof)
    if not hdr:
        return None
    probe = Binding(prof["profile"])
    _bind_columns(ws, hdr, prof, probe)
    required = [f for f in prof["required"] if probe.columns.get(f)]
    missing = [f for f in prof["required"] if not probe.columns.get(f)]

    from_first_digit = prof.get("isced_from_first_digit_of")
    col_level = probe.columns.get(from_first_digit or "isced_code")
    col_label = probe.columns.get("isced_label")
    col_prog = probe.columns.get("programme_en")
    parsed = 0
    seen = 0
    for r in range(hdr + 1, min((ws.max_row or 0), hdr + 300) + 1):
        prog = _txt(ws, r, col_prog) if col_prog else ""
        raw = _txt(ws, r, col_level) if col_level else ""
        lab = _txt(ws, r, col_label) if col_label else ""
        if not prog and not raw:
            continue
        seen += 1
        code = (P.three_digit(raw)[:1] if from_first_digit
                else P.isced_code(raw, lab))
        if code in ISCED_LEVELS and prog:
            parsed += 1

    score = (0 if missing else 1000) + parsed * 5 + len(probe.headers)
    if _looks_like_prose(ws, hdr):
        score -= 500
    return {
        "profile": prof["profile"], "sheet": name, "header_row": hdr,
        "bound": len(probe.headers), "required_bound": len(required),
        "required_missing": missing,
        "rows_seen": seen, "rows_parsed": parsed, "score": score,
        "prose": _looks_like_prose(ws, hdr),
    }


def _find_header(ws, prof) -> Optional[int]:
    anchors = {P.norm(a) for a in prof["header"]["anchor"]}
    wanted = []
    for cands in prof["columns"].values():
        wanted += [P.norm(c) for c in cands if not c.startswith("__")]
    wanted = set(wanted)
    limit = min(ws.max_row or 1, prof["header"]["search_rows"])
    for r in range(1, limit + 1):
        if P.norm(_txt(ws, r, 1)) not in anchors:
            continue
        hits = 0
        for c in range(1, (ws.max_column or 1) + 1):
            h = P.norm(_txt(ws, r, c))
            if not h:
                continue
            if any(h.startswith(w) or w.startswith(h) for w in wanted if w):
                hits += 1
        if hits >= prof["header"]["min_hits"]:
            return r
    return None


def _bind_columns(ws, header_row: int, prof, b: Binding) -> None:
    ncol = ws.max_column or 1
    headers = {c: _txt(ws, header_row, c) for c in range(1, ncol + 1)}
    nheaders = {c: P.norm(h) for c, h in headers.items()}
    used = set()

    # Pass 1: exact header matches, so a loose candidate on one field can never
    # take a column another field names exactly.
    for fieldname, cands in prof["columns"].items():
        if any(str(c).startswith("__") for c in cands):
            continue
        for cand in cands:
            n = P.norm(cand)
            hit = next((c for c, h in nheaders.items()
                        if h and h == n and c not in used), None)
            if hit:
                used.add(hit)
                b.columns[fieldname] = hit
                b.headers[fieldname] = headers[hit]
                break

    for fieldname, cands in prof["columns"].items():
        if b.columns.get(fieldname):
            continue
        if any(str(c).startswith("__") for c in cands):
            continue
        best = None
        for cand in cands:                       # profile order is priority
            n = P.norm(cand)
            for c, h in nheaders.items():
                if c in used or not h:
                    continue
                if h == n:
                    best = c
                    break
            if best:
                break
        if best is None:
            for cand in cands:
                n = P.norm(cand)
                for c, h in nheaders.items():
                    if c in used or not h:
                        continue
                    if h.startswith(n) or n.startswith(h):
                        best = c
                        break
                if best:
                    break
        if best:
            used.add(best)
            b.columns[fieldname] = best
            b.headers[fieldname] = headers[best]
        else:
            b.columns[fieldname] = None
            b.unbound.append(fieldname)
            b.candidates[fieldname] = [
                (c, headers[c]) for c in sorted(headers)
                if headers[c] and c not in used][:6]

    # A profile that derives the level from a 3-digit code has no level column
    # and should not be reported as missing one.
    if prof.get("isced_from_first_digit_of"):
        b.columns.setdefault("isced_code", None)
        return

    # the ISCED level *code* sits in the unlabelled column beside its label
    lab = b.columns.get("isced_label")
    if lab:
        for c in range(lab + 1, min(lab + 3, ncol + 1)):
            if not headers.get(c):
                b.columns["isced_code"] = c
                b.headers["isced_code"] = (
                    f"(unlabelled, right of {headers[lab]!r})")
                break
    if not b.columns.get("isced_code"):
        b.columns["isced_code"] = None
        b.unbound.append("isced_code")


def _identity(ws, path, prof, b: Binding,
              header_row: Optional[int] = None) -> Tuple[str, str, Optional[int]]:
    """(iso3, country, reference year).

    The banner is written in the workbook's own language -- *Cartographie de
    la CITE 2011 de Sénégal* -- so the country is found by looking up every
    contiguous run of words against the ISO 3166 table rather than by
    stripping a known English phrase. A profile whose country lives in a
    column says so with ``country_column``.
    """
    name = ""
    ref = None
    banner = prof.get("banner") or {}
    for r in range(1, (banner.get("title_row_max") or 0) + 1):
        for c in range(1, 4):
            t = _txt(ws, r, c)
            if not t:
                continue
            for pat in banner.get("title_patterns") or []:
                m = re.search(pat, t, re.I)
                if m and not name:
                    name = m.group("name").strip()
            for pat in banner.get("reference_patterns") or []:
                m = re.search(pat, t, re.I)
                if m and ref is None:
                    ref = int(m.group("year"))
            if not name:
                got, span = CO.match(t)
                if got:
                    name = span
    # a profile that carries the country and the year in columns
    col = banner.get("country_column")
    if col and header_row and b.columns.get(col):
        for r in range(header_row + 1, header_row + 40):
            v = _txt(ws, r, b.columns[col])
            if v:
                name = name or v
                break
    if b.columns.get("school_year") and ref is None and header_row:
        for r in range(header_row + 1, header_row + 40):
            v = _txt(ws, r, b.columns["school_year"])
            if not v:
                continue
            for pat in banner.get("reference_patterns") or []:
                m = re.search(pat, v, re.I)
                if m:
                    ref = int(m.group("year"))
                    break
            if ref:
                break

    iso3, name = CO.resolve(name, str(path))
    b.identity = {"country": name, "iso3": iso3, "reference_year": ref,
                  "country_en": CO.english_name(iso3)}
    return iso3, name, ref


def _language(inst: Instruction, prof, b: Binding) -> Dict[str, Any]:
    """What the workbook holds in the country's own language, and which
    language that is likely to be.

    The mapping workbook pairs every substantive column with a national-language
    twin.  Those twins are the only record of what a programme is actually
    called where it is taught -- "Trung h\u1ecdc c\u01a1 s\u1edf", not "Lower secondary" -- and
    once they are dropped they cannot be recovered from anything downstream.
    They are kept verbatim, per row.

    The workbook never states its language, so the language *name* is inferred
    from the country and carries that as its evidence.  A country with no entry
    gets no hint rather than a guess.
    """
    cfg = prof.get("language") or {}
    cols = cfg.get("national_columns") or []
    present = [c for c in cols if b.columns.get(c)]
    filled = {c: 0 for c in present}
    non_ascii = 0
    for r in inst.rows:
        for c in present:
            fld = {"programme_national": r.programme_national,
                   "entrance_req_national": r.entrance_requirement_national,
                   "diploma_national": r.diploma_national}.get(c, "")
            if fld and not P.is_na(fld, prof["na_tokens"]):
                filled[c] += 1
                if any(ord(ch) > 127 for ch in fld):
                    non_ascii += 1
    hint = (cfg.get("by_iso3") or {}).get(inst.iso3)
    names = [hint] if isinstance(hint, str) else list(hint or [])
    out: Dict[str, Any] = {
        "national_columns": {c: {"header": b.headers.get(c, ""),
                                 "rows_with_text": filled[c]}
                             for c in present},
        "rows_with_national_text": max(filled.values()) if filled else 0,
        "non_latin_or_accented_values": non_ascii,
        "english_columns_present": [c for c in (cfg.get("english_columns") or [])
                                    if b.columns.get(c)],
    }
    if names:
        out["national_language"] = inferred(
            names, 0.6,
            f"not stated in the workbook; inferred from the country "
            f"({inst.iso3} {inst.country}) via the profile's by_iso3 map")
    else:
        out["national_language"] = None
        out["note"] = (f"no national-language hint for {inst.iso3 or '??'}; "
                       f"the national text is kept verbatim and unattributed")
    return out


# --------------------------------------------------------------------------
def read(path, iso3: Optional[str] = None, country: Optional[str] = None,
         version: str = "v1.0", profile_path: Optional[str] = None,
         reference_year: Optional[int] = None,
         sheet: Optional[str] = None) -> Tuple[Instruction, Binding]:
    """Read one ISCED mapping workbook, whichever family and sheet it is.

    Every sheet is scored against every profile and the best pair is read.
    ``sheet=`` forces one when a reviewer knows better than the score.
    """
    wb = load_workbook(path, data_only=True, read_only=False)
    profiles = load_profiles(profile_path)

    candidates: List[dict] = []
    for prof in profiles:
        names = [s for s in prof["sheet"]["candidates"] if s in wb.sheetnames]
        if prof["sheet"].get("any_sheet"):
            names += [s for s in wb.sheetnames if s not in names]
        for name in names:
            if sheet and name != sheet:
                continue
            try:
                got = score_sheet(wb[name], prof, name)
            except Exception:                              # noqa: BLE001
                got = None
            if got:
                got["priority"] = prof.get("priority", 100)
                candidates.append(got)

    # required columns first, then how much actually parsed, then the more
    # specific profile
    candidates.sort(key=lambda c: (-(not c["required_missing"]), -c["score"],
                                   c["priority"]))
    if not candidates or candidates[0]["required_missing"]:
        inst = Instruction(version=version)
        b = Binding(profiles[0]["profile"] if profiles else "")
        b.candidates_considered = candidates
        b.tried = [f"{c['profile']} on {c['sheet']!r}: "
                   f"{c['rows_parsed']} row(s) parsed"
                   + (f", missing {', '.join(c['required_missing'])}"
                      if c["required_missing"] else "")
                   for c in candidates]
        inst.source = SourceFile.of(path, "", None, 0,
                                    b.profile)
        b.problems.append(
            "no sheet in this workbook reads as an ISCED mapping. "
            + ("Considered: " + "; ".join(b.tried)
               if b.tried else
               "No sheet carried a recognisable header row.")
            + " If it is a new template family, add a profile YAML beside the "
              "others; if the table is on a sheet the score missed, pass "
              "--sheet.")
        return inst, b

    best = candidates[0]
    prof = next(p for p in profiles if p["profile"] == best["profile"])
    inst, b = _read_with(wb, path, prof, iso3, country, version,
                         reference_year, sheet_name=best["sheet"])
    b.candidates_considered = candidates
    b.score = best
    b.tried = [f"{c['profile']} on {c['sheet']!r}: {c['rows_parsed']} parsed, "
               f"score {c['score']}" for c in candidates]
    runner_up = candidates[1] if len(candidates) > 1 else None
    if runner_up and not runner_up["required_missing"] and \
            runner_up["score"] >= best["score"] * 0.9:
        b.problems.append(
            f"two sheets read almost equally well as an ISCED mapping: "
            f"{best['sheet']!r} ({best['rows_parsed']} rows) was taken over "
            f"{runner_up['sheet']!r} ({runner_up['rows_parsed']} rows). "
            f"Check that is the right one, or pass --sheet.")
    return inst, b


def _read_with(wb, path, prof, iso3, country, version, reference_year,
               sheet_name: Optional[str] = None) -> Tuple[Instruction, Binding]:
    b = Binding(prof["profile"])
    na = prof["na_tokens"]

    order = ([sheet_name] if sheet_name else
             [s for s in prof["sheet"]["candidates"] if s in wb.sheetnames]
             + ([s for s in wb.sheetnames
                 if s not in prof["sheet"]["candidates"]]
                if prof["sheet"].get("any_sheet") else []))
    ws = hdr = None
    for name in order:
        if name not in wb.sheetnames:
            continue
        cand = wb[name]
        h = _find_header(cand, prof)
        if h:
            ws, hdr = cand, h
            b.sheet = name
            b.header_row = h
            break
    inst = Instruction(version=version, frame=prof["frame"])
    if ws is None:
        inst.source = SourceFile.of(path, "", None, 0, prof["profile"])
        return inst, b

    _bind_columns(ws, hdr, prof, b)
    fiso, fname, ref = _identity(ws, path, prof, b, hdr)
    ref = reference_year or ref
    inst.iso3 = (iso3 or fiso or "").upper()
    inst.country = country or fname
    if not inst.iso3:
        b.problems.append(
            f"the country could not be determined for "
            f"{fname or '(unnamed)'}; pass --iso3")
    if ref is None:
        b.problems.append(
            "no school-year reference found; the era window cannot be "
            "placed -- pass --reference-year")

    missing = [f for f in prof["required"] if not b.columns.get(f)]
    if missing:
        b.problems.append(
            f"required field(s) did not bind: {', '.join(missing)}. This is a "
            f"template change; add the real header string to "
            f"{prof['profile']}.yaml.")
        return inst, b

    col = b.columns
    gmd_map = prof["value_maps"]["gmd_from_isced"]
    comp_map = prof["value_maps"]["completion"]
    orient_map = prof["value_maps"]["orientation"]
    from_first_digit = prof.get("isced_from_first_digit_of")
    keep = prof.get("row_filter")

    def cell(r, field):
        c = col.get(field)
        return _txt(ws, r, c) if c else ""

    rows: List[EduRow] = []
    r = hdr + 1
    blanks = 0
    seq = 0
    while r <= (ws.max_row or 0) and blanks < 6:
        prog = cell(r, "programme_en")
        pid_raw = _txt(ws, r, 1)
        has_id = bool(re.fullmatch(r"\d+", pid_raw or ""))
        if not prog and not has_id:
            blanks += 1
            r += 1
            continue
        if not prog:
            blanks += 1
            r += 1
            continue
        blanks = 0
        if keep and col.get(keep["column"]):
            v = P.norm(cell(r, keep["column"]))
            if v and v not in {P.norm(k) for k in keep["keep"]}:
                r += 1
                continue

        if from_first_digit:
            code = P.three_digit(cell(r, from_first_digit))
            isced = code[0] if code else ""
        else:
            isced = P.isced_code(cell(r, "isced_code"), cell(r, "isced_label"))
        if isced not in ISCED_LEVELS:
            inst.unmapped.append(Unmapped(
                "rows.isced", r, cell(r, from_first_digit or "isced_code"),
                "not an ISCED 2011 level code"))
            r += 1
            continue

        acc = P.access(cell(r, "access"), na,
                       prof["value_maps"].get("access"))
        comp = P.completion(cell(r, "completion"), comp_map, na)
        if not col.get("completion") and col.get("position"):
            comp = {"completion": "full" if cell(r, "position") else "na",
                    "position": cell(r, "position")}
        seq += 1
        rows.append(EduRow(
            era=1,
            programme=prog,
            programme_national=cell(r, "programme_national"),
            isced=isced,
            isced_label=ISCED_LEVELS[isced],
            isced_p=P.three_digit(cell(r, "isced_p")),
            isced_a=P.three_digit(cell(r, "isced_a")),
            orientation=P.orientation(cell(r, "orientation"), orient_map, na),
            completion=comp["completion"],
            position=comp["position"],
            access=acc["access"],
            access_to=acc["to"],
            entrance_age=P.entrance_age(cell(r, "entrance_age"), na),
            duration_years=P.duration(cell(r, "duration"), na),
            complete_at=cell(r, "diploma_en"),
            gmd=gmd_map.get(isced, ""),
            entrance_requirement=cell(r, "entrance_req_en"),
            entrance_requirement_national=cell(r, "entrance_req_national"),
            diploma=cell(r, "diploma_en"),
            diploma_national=cell(r, "diploma_national"),
            notes=cell(r, "notes"),
            programme_id=int(pid_raw) if has_id else seq,
            source_row=r,
        ))
        r += 1

    b.rows_read = len(rows)
    inst.rows = rows
    inst.language = _language(inst, prof, b)
    inst.source = SourceFile.of(path, b.sheet, hdr, len(rows), prof["profile"],
                                reference_year=ref)
    return inst, b
