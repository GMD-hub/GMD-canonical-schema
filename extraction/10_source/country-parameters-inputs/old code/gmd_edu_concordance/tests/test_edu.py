"""
Tests for gmd_edu_concordance.

The workbook-dependent tests are skipped unless an ISCED mapping workbook is
present.  Point GMD_ISCED_DIR at a directory of them, or drop them in
./samples/.

    python3 tests/test_edu.py
    GMD_ISCED_DIR=/path/to/files python3 tests/test_edu.py
"""
import glob
import shutil
import urllib.error
import urllib.request
import json
import os
import sys
import tempfile
import unittest

import yaml


def _walk(d):
    out = []
    for root, _dirs, files in os.walk(d):
        out += [os.path.join(root, f) for f in files]
    return out

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from gmd_edu_concordance import check, extract, resolve             # noqa: E402
from gmd_edu_concordance import countries as CO
from gmd_edu_concordance import registry as REG
from gmd_edu_concordance import review as RV
from gmd_edu_concordance import viewer as V
from gmd_edu_concordance.viewer import build as build_viewer        # noqa: E402
from gmd_edu_concordance.derive import run as derive_run            # noqa: E402
from gmd_edu_concordance.derive import consistency                  # noqa: E402
from gmd_edu_concordance.reader import score_sheet                  # noqa: E402
from gmd_edu_concordance import parse as P                          # noqa: E402
from gmd_edu_concordance.reader import (                            # noqa: E402
    load_profile, load_profiles, read,
)
from gmd_edu_concordance.schema import (                            # noqa: E402
    GMD_EDUC_ATTAINMENT, GMD_EDUC_LEVELS, ISCED_LEVELS, attainment_of,
    is_inferred, value_of,
)
from gmd_edu_concordance.validate import validate as run_rules      # noqa: E402

SAMPLE_DIRS = [os.environ.get("GMD_ISCED_DIR"),
               os.path.join(os.path.dirname(__file__), "..", "samples"),
               os.path.dirname(__file__)]

NA = ["-", "n/a", "na", "none", "."]


def find_workbooks():
    for d in SAMPLE_DIRS:
        if not d or not os.path.isdir(d):
            continue
        hits = sorted(glob.glob(os.path.join(d, "ISCED_*.xlsx")))
        if hits:
            return hits
    return []


WB = find_workbooks()
needs_wb = unittest.skipUnless(WB, "no ISCED workbook found; set GMD_ISCED_DIR")


class TestParse(unittest.TestCase):
    def test_duration_single_and_range(self):
        self.assertEqual(P.duration("6", NA)["min"], 6)
        d = P.duration("3-4", NA)
        self.assertTrue(d["range"])
        self.assertEqual((d["min"], d["max"]), (3, 4))

    def test_months_become_fractions_of_a_year(self):
        self.assertEqual(P.duration("6 Months", NA)["min"], 0.5)
        self.assertEqual(P.entrance_age("3 months", NA)["min"], 0.25)

    def test_na_is_none_not_zero(self):
        for v in ("-", "n/a", "", None):
            self.assertIsNone(P.duration(v, NA))

    def test_isced_code_from_label_or_code(self):
        self.assertEqual(P.isced_code(1), "1")
        self.assertEqual(P.isced_code("", "ISCED 3 Upper secondary"), "3")
        self.assertEqual(P.isced_code("nonsense"), "")

    def test_completion_splits_position(self):
        m = {"full": ["full completion"], "partial": ["partial completion"],
             "none": ["no completion"]}
        got = P.completion("Full completion: First degree (3-4 years)", m, NA)
        self.assertEqual(got["completion"], "full")
        self.assertEqual(got["position"], "First degree (3-4 years)")

    def test_access_keeps_the_target_level(self):
        self.assertEqual(P.access("Yes, to ISCED 3", NA),
                         {"access": True, "to": "ISCED 3"})
        self.assertEqual(P.access("No", NA)["access"], False)
        self.assertIsNone(P.access("-", NA)["access"])

    def test_reform_signals(self):
        got = P.reform_signals("Reform in 2013 changed the cycle.")
        self.assertEqual(got[0]["year"], 2013)


class TestProfile(unittest.TestCase):
    def test_profile_lists_required_fields(self):
        p = load_profile()
        for f in ("programme_en", "isced_code", "entrance_age", "duration"):
            self.assertIn(f, p["required"])

    def test_gmd_map_targets_are_controlled(self):
        p = load_profile()
        for v in p["value_maps"]["gmd_from_isced"].values():
            self.assertIn(v, GMD_EDUC_LEVELS)

    def test_every_isced_level_has_a_gmd_target(self):
        p = load_profile()
        for lvl in ISCED_LEVELS:
            self.assertIn(lvl, p["value_maps"]["gmd_from_isced"])


@needs_wb
class TestRealWorkbooks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cases = []
        for f in WB:
            inst, b = read(f)
            derive_run(inst, load_profile())
            cls.cases.append((f, inst, b))

    def test_identity_reference_year_and_rows(self):
        for f, inst, b in self.cases:
            self.assertEqual(len(inst.iso3), 3, f)
            self.assertTrue(inst.source.reference_year, f)
            self.assertGreater(len(inst.rows), 5, f)
            self.assertEqual(b.problems, [], f)

    def test_required_columns_bind(self):
        profs = {p["profile"]: p for p in load_profiles()}
        for f, inst, b in self.cases:
            for fld in profs[b.profile]["required"]:
                self.assertIsNotNone(b.columns.get(fld), f"{f}: {fld}")

    def test_grades_are_derived_and_marked_inferred(self):
        for f, inst, b in self.cases:
            prim = [r for r in inst.rows if r.isced == "1"][0]
            self.assertTrue(is_inferred(prim.grades), f)
            self.assertTrue(value_of(prim.grades).startswith("1-"), f)
            self.assertIn("theoretical", prim.grades["evidence"])

    def test_grade_1_starts_at_the_isced_1_entrance_age(self):
        for f, inst, b in self.cases:
            origin = inst.structure["grade_origin"]
            prim = [r for r in inst.rows if r.isced == "1"][0]
            self.assertEqual(origin["age"], prim.entrance_age["min"], f)
            self.assertEqual(value_of(prim.grade_from), 1, f)

    def test_years_in_level_comes_from_the_duration_column(self):
        for f, inst, b in self.cases:
            for r in inst.rows:
                if r.duration_years:
                    self.assertEqual(
                        value_of(r.years_in_level),
                        P.fmt_range(r.duration_years["min"],
                                    r.duration_years["max"]), f)

    def test_the_ladder_has_no_holes(self):
        for f, inst, b in self.cases:
            gs = [s.grade for s in inst.ladder]
            self.assertEqual(gs, list(range(min(gs), max(gs) + 1)), f)

    def test_tertiary_carries_years_but_never_a_grade(self):
        for f, inst, b in self.cases:
            for r in inst.rows:
                if r.isced in ("5", "6", "7", "8"):
                    self.assertEqual(value_of(r.grades), "-", f)
                    self.assertIsNotNone(r.years_at_completion, f)

    def test_the_workbook_is_the_latest_era(self):
        for f, inst, b in self.cases:
            self.assertEqual(len(inst.eras), 1, f)
            era = inst.eras[0]
            self.assertTrue(era.latest, f)
            self.assertEqual(value_of(era.to_year), 2035, f)
            self.assertTrue(is_inferred(era.from_year), f)

    def test_a_cohort_before_the_era_does_not_default_into_it(self):
        for f, inst, b in self.cases:
            start = value_of(inst.eras[0].from_year)
            got = resolve(inst, grade=6, year_last_in_school=start - 30)
            self.assertFalse(got["resolved"], f)
            self.assertIn("ED-01", got["why"])

    def test_a_cohort_inside_the_era_resolves(self):
        for f, inst, b in self.cases:
            y = inst.source.reference_year
            got = resolve(inst, grade=1, year_last_in_school=y)
            self.assertTrue(got["resolved"], f)
            self.assertEqual(got["gmd"], "primary", f)

    def test_completion_is_true_only_on_the_last_grade_of_a_level(self):
        for f, inst, b in self.cases:
            prim = [s for s in inst.ladder if s.gmd == "primary"]
            self.assertTrue(prim[-1].completes_level, f)
            self.assertFalse(prim[0].completes_level, f)

    def test_every_row_carries_a_controlled_gmd_target(self):
        for f, inst, b in self.cases:
            for r in inst.rows:
                self.assertIn(r.gmd, GMD_EDUC_LEVELS, f)

    def test_e03_always_fires_for_a_single_vintage_workbook(self):
        for f, inst, b in self.cases:
            rules = {x.rule for x in run_rules(inst, b)}
            self.assertIn("E-03", rules, f)

    def test_no_blocking_findings_on_a_well_formed_workbook(self):
        for f, inst, b in self.cases:
            blocks = [x for x in run_rules(inst, b) if x.level == "BLOCK"]
            self.assertEqual(blocks, [], f"{f}: {[x.message for x in blocks]}")

    def test_fingerprint_is_stable(self):
        f, inst, _ = self.cases[0]
        again, _ = read(f)
        derive_run(again, load_profile())
        self.assertEqual(inst.content_fingerprint(),
                         again.content_fingerprint())

    def test_extract_writes_seven_artifacts(self):
        with tempfile.TemporaryDirectory() as d:
            res = extract(WB[0], out=d)
            # 7 artefacts plus the shared registry
            self.assertEqual(len(res["artifacts"]), 8)
            for p in res["artifacts"]:
                self.assertTrue(os.path.getsize(p) > 0, p)
            names = {os.path.basename(p) for p in res["artifacts"]}
            iso = res["instruction"].iso3
            self.assertIn(f"{iso}_edu_view.json", names)
            self.assertIn(f"{iso}_edu_ladder.json", names)

    # -- national language ------------------------------------------------
    def test_national_language_columns_are_kept(self):
        for f, inst, b in self.cases:
            lang = inst.language
            self.assertIn("programme_national", lang["national_columns"], f)
            self.assertGreater(lang["rows_with_national_text"], 0, f)

    def test_national_programme_text_survives_to_the_row(self):
        for f, inst, b in self.cases:
            self.assertTrue(all(r.programme_national for r in inst.rows), f)

    def test_a_non_latin_workbook_keeps_its_script(self):
        vnm = [c for c in self.cases if c[1].iso3 == "VNM"]
        if not vnm:
            self.skipTest("Viet Nam workbook not available")
        inst = vnm[0][1]
        self.assertGreater(inst.language["non_latin_or_accented_values"], 0)
        self.assertTrue(any(any(ord(ch) > 127 for ch in r.programme_national)
                            for r in inst.rows))

    def test_language_name_is_inferred_with_evidence_never_asserted(self):
        for f, inst, b in self.cases:
            nl = inst.language.get("national_language")
            if nl is None:
                continue
            self.assertTrue(is_inferred(nl), f)
            self.assertIn("not stated in the workbook", nl["evidence"])

    # -- standalone validator ---------------------------------------------
    def test_standalone_check_reproduces_the_extraction_verdict(self):
        with tempfile.TemporaryDirectory() as d:
            res = extract(WB[0], out=d)
            again = check(d)
            self.assertEqual(len(again), 1)
            self.assertEqual(again[0]["verdict"], res["verdict"])
            self.assertEqual(again[0]["counts"], res["counts"])

    def test_the_check_catches_an_edit_that_breaks_a_rule(self):
        """A hand edit that puts a hole in the grade ladder must be caught
        without the workbook."""
        with tempfile.TemporaryDirectory() as d:
            extract(WB[0], out=d)
            path = [p for p in _walk(d) if p.endswith("_edu_v1.0.yaml")][0]
            with open(path, encoding="utf-8") as fh:
                doc = yaml.safe_load(fh)
            doc["ladder"] = [s for s in doc["ladder"] if s["grade"] != 3]
            with open(path, "w", encoding="utf-8") as fh:
                yaml.safe_dump(doc, fh, sort_keys=False, allow_unicode=True)
            res = check(d)[0]
            self.assertEqual(res["verdict"], "blocked")
            self.assertTrue(any(x.rule == "E-10" and x.level == "BLOCK"
                                for x in res["findings"]))

    # -- viewer -----------------------------------------------------------
    def test_viewer_is_one_self_contained_file(self):
        with tempfile.TemporaryDirectory() as d:
            extract(WB[0], out=d)
            html = build_viewer(d, os.path.join(d, "v.html"))
            with open(html, encoding="utf-8") as fh:
                body = fh.read()
            self.assertGreater(os.path.getsize(html), 20000)
            self.assertIn("__BUNDLES__", body)
            for bad in ("http://", "https://", "<script src", "<link rel"):
                self.assertNotIn(bad, body, f"viewer reaches for {bad}")

    def test_the_view_bundle_carries_what_the_viewer_needs(self):
        with tempfile.TemporaryDirectory() as d:
            res = extract(WB[0], out=d)
            path = [p for p in res["artifacts"] if p.endswith("_view.json")][0]
            with open(path, encoding="utf-8") as fh:
                v = json.load(fh)
            for k in ("kind", "eras", "rows", "ladder", "findings", "counts",
                      "structure", "coverage", "language", "probe"):
                self.assertIn(k, v)
            self.assertEqual(v["kind"], "education")
            self.assertTrue(all("grades_v" in r for r in v["rows"]))



# ==========================================================================
# the review app
# ==========================================================================
MINIMAL_BUNDLE = {
    "kind": "wash", "domain": "water",
    "country": {"iso3": "ZZZ", "name": "Testland"},
    "version": "v1.0", "status": "draft", "verdict": "clean",
    "counts": {"BLOCK": 0, "WARN": 0, "INFO": 1},
    "rows": [], "classifications": [], "facility_types": [], "sources": [],
    "findings": [], "source": {"file": "x.xlsx"},
}


def _fake_out(root):
    """An output tree with no workbook behind it."""
    d = os.path.join(root, "ZZZ")
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "ZZZ_water_view.json"), "w") as fh:
        json.dump(MINIMAL_BUNDLE, fh)
    with open(os.path.join(d, "ZZZ_water_v1.0.yaml"), "w") as fh:
        fh.write("schema: gmd.country_instruction.wash.water\n")
    with open(os.path.join(d, "ZZZ_secret.env"), "w") as fh:
        fh.write("TOKEN=hunter2\n")
    return d


class TestViewerPage(unittest.TestCase):
    def test_picker_page_carries_no_data(self):
        with tempfile.TemporaryDirectory() as d:
            p = V.build(None, os.path.join(d, "v.html"))
            body = open(p, encoding="utf-8").read()
            self.assertIn('window.__MODE__ = "picker"', body)
            self.assertIn("window.__BUNDLES__ = [];", body)

    def test_embed_page_carries_the_bundles(self):
        with tempfile.TemporaryDirectory() as d:
            _fake_out(d)
            p = V.build(d, os.path.join(d, "v.html"), embed=True)
            body = open(p, encoding="utf-8").read()
            self.assertIn('window.__MODE__ = "embed"', body)
            self.assertIn("Testland", body)

    def test_mode_placeholder_does_not_eat_the_variable_name(self):
        """`__MODE__` used to substitute into `window.__MODE__` itself."""
        with tempfile.TemporaryDirectory() as d:
            body = open(V.build(None, os.path.join(d, "v.html")),
                        encoding="utf-8").read()
            self.assertIn("window.__MODE__", body)

    def test_page_never_reaches_for_the_network(self):
        with tempfile.TemporaryDirectory() as d:
            _fake_out(d)
            for embed in (False, True):
                body = open(V.build(d, os.path.join(d, "v.html"), embed=embed),
                            encoding="utf-8").read()
                for bad in ("http://", "https://", "<script src", "<link rel"):
                    self.assertNotIn(bad, body, f"embed={embed} pulls {bad}")


class TestViewerScan(unittest.TestCase):
    def test_scan_finds_a_country_folder(self):
        with tempfile.TemporaryDirectory() as d:
            _fake_out(d)
            sc = V.scan(d)
            self.assertEqual(len(sc["countries"]), 1)
            c = sc["countries"][0]
            self.assertEqual((c["iso3"], c["dir"]), ("ZZZ", "ZZZ"))
            self.assertIn("water", [x["domain"] for x in c["domains"]])

    def test_two_folders_for_one_country_are_two_entries(self):
        """A country is a folder. Keying by ISO3 silently dropped the second."""
        with tempfile.TemporaryDirectory() as d:
            _fake_out(d)
            shutil.copytree(os.path.join(d, "ZZZ"), os.path.join(d, "ZZZ_v2"))
            self.assertEqual(len(V.scan(d)["countries"]), 2)

    def test_scan_lists_only_readable_artefacts(self):
        with tempfile.TemporaryDirectory() as d:
            _fake_out(d)
            names = {f["name"] for f in V.scan(d)["countries"][0]["files"]}
            self.assertIn("ZZZ_water_v1.0.yaml", names)
            self.assertNotIn("ZZZ_secret.env", names)

    def test_scan_is_reread_not_cached(self):
        with tempfile.TemporaryDirectory() as d:
            _fake_out(d)
            self.assertEqual(len(V.scan(d)["countries"]), 1)
            shutil.copytree(os.path.join(d, "ZZZ"), os.path.join(d, "YYY"))
            self.assertEqual(len(V.scan(d)["countries"]), 2)


class TestViewerPaths(unittest.TestCase):
    def test_traversal_is_refused_not_normalised(self):
        with tempfile.TemporaryDirectory() as d:
            root = os.path.abspath(d)
            for bad in ("../x", "../../etc/passwd", "/etc/passwd", "a/../../b",
                        "..", "ZZZ/../../x"):
                self.assertIsNone(V._safe_join(root, bad), bad)

    def test_ordinary_paths_resolve(self):
        with tempfile.TemporaryDirectory() as d:
            root = os.path.abspath(d)
            self.assertEqual(V._safe_join(root, "ZZZ/a.yaml"),
                             os.path.join(root, "ZZZ", "a.yaml"))
            self.assertEqual(V._safe_join(root, ""), root)


class TestViewerServer(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.mkdtemp()
        _fake_out(cls.dir)
        cls.httpd, cls.url = V.serve(cls.dir, port=0, block=False)

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        shutil.rmtree(cls.dir, ignore_errors=True)

    def get(self, path):
        with urllib.request.urlopen(self.url.rstrip("/") + path, timeout=10) as r:
            return r.read().decode()

    def test_index_is_the_page_in_server_mode(self):
        body = self.get("/")
        self.assertIn('window.__MODE__ = "server"', body)
        self.assertIn("window.__BUNDLES__ = [];", body)

    def test_scan_and_bundles(self):
        sc = json.loads(self.get("/api/scan"))
        self.assertEqual(sc["countries"][0]["iso3"], "ZZZ")
        b = json.loads(self.get("/api/bundles?dir=ZZZ"))
        self.assertEqual(b[0]["country"]["iso3"], "ZZZ")

    def test_empty_dir_means_every_folder(self):
        b = json.loads(self.get("/api/bundles?dir="))
        self.assertTrue(b)

    def test_raw_reads_an_artefact(self):
        d = json.loads(self.get("/api/raw?rel=ZZZ/ZZZ_water_v1.0.yaml"))
        self.assertIn("gmd.country_instruction.wash.water", d["text"])

    def test_a_new_folder_appears_without_a_restart(self):
        """The entire point: rerun the extractor, hit reload, see it."""
        before = len(json.loads(self.get("/api/scan"))["countries"])
        shutil.copytree(os.path.join(self.dir, "ZZZ"),
                        os.path.join(self.dir, "QQQ"))
        try:
            after = len(json.loads(self.get("/api/scan"))["countries"])
            self.assertEqual(after, before + 1)
        finally:
            shutil.rmtree(os.path.join(self.dir, "QQQ"))

    def test_it_refuses_to_leave_the_served_root(self):
        for bad in ("/api/raw?rel=../../etc/passwd", "/api/raw?rel=/etc/passwd",
                    "/api/bundles?dir=../..", "/api/bundles?dir=/tmp"):
            with self.assertRaises(urllib.error.HTTPError, msg=bad):
                self.get(bad)

    def test_it_refuses_extensions_that_are_not_artefacts(self):
        with self.assertRaises(urllib.error.HTTPError):
            self.get("/api/raw?rel=ZZZ/ZZZ_secret.env")

    def test_it_binds_loopback_only(self):
        self.assertEqual(self.httpd.server_address[0], "127.0.0.1")


# ==========================================================================
# language and template family
# ==========================================================================
class TestCountries(unittest.TestCase):
    def test_a_country_named_in_its_own_language(self):
        for name, iso in (("España", "ESP"), ("Sénégal", "SEN"),
                          ("Sri Lanka", "LKA"), ("Viet Nam", "VNM")):
            self.assertEqual(CO.to_iso3(name), iso, name)

    def test_a_country_named_inside_a_banner(self):
        for banner, iso in (
                ("Sri Lanka - ISCED 2011 Mapping", "LKA"),
                ("Argentina ISCED 2011 Mapping", "ARG"),
                ("Cartographie de la CITE 2011 de Sénégal", "SEN")):
            self.assertEqual(CO.resolve(banner)[0], iso, banner)

    def test_an_unknown_name_is_none_never_a_guess(self):
        self.assertIsNone(CO.to_iso3("Nowhereland"))


class TestProfiles(unittest.TestCase):
    def test_more_than_one_profile_ships(self):
        names = {p["profile"] for p in load_profiles()}
        self.assertIn("isced_uis_2011_mapping", names)
        self.assertIn("uoe_scope_mapping", names)

    def test_profiles_are_ordered_most_specific_first(self):
        ps = load_profiles()
        self.assertEqual([p.get("priority", 100) for p in ps],
                         sorted(p.get("priority", 100) for p in ps))

    def test_every_profile_declares_what_it_needs(self):
        for p in load_profiles():
            for key in ("profile", "frame", "sheet", "header", "columns",
                        "required", "value_maps", "na_tokens"):
                self.assertIn(key, p, p.get("profile"))
            for f in p["required"]:
                self.assertIn(f, p["columns"], p["profile"])

    def test_the_iso3166_table_is_not_a_profile(self):
        self.assertNotIn("iso3166", {p["profile"] for p in load_profiles()})


@needs_wb
class TestTranslatedWorkbooks(unittest.TestCase):
    """A workbook in French, or on another template, must read like the rest."""

    @classmethod
    def setUpClass(cls):
        cls.by_iso = {}
        profs = {p["profile"]: p for p in load_profiles()}
        for f in WB:
            inst, b = read(f)
            if b.profile in profs:
                derive_run(inst, profs[b.profile])
            cls.by_iso[inst.iso3 or os.path.basename(f)] = (inst, b, f)

    def test_every_workbook_finds_its_country(self):
        for k, (inst, b, f) in self.by_iso.items():
            self.assertEqual(len(inst.iso3), 3, os.path.basename(f))

    def test_every_workbook_binds_a_profile_and_reads_rows(self):
        for k, (inst, b, f) in self.by_iso.items():
            self.assertTrue(b.header_row, k)
            self.assertTrue(inst.rows, k)
            self.assertTrue(inst.source.reference_year, k)

    def test_required_columns_bind_in_every_language(self):
        profs = {p["profile"]: p for p in load_profiles()}
        for k, (inst, b, f) in self.by_iso.items():
            for fld in profs[b.profile]["required"]:
                self.assertIsNotNone(b.columns.get(fld), f"{k}: {fld}")

    def test_a_french_workbook_reads_its_french_values(self):
        sen = self.by_iso.get("SEN")
        if not sen:
            self.skipTest("Senegal workbook not available")
        inst = sen[0]
        prim = [r for r in inst.rows if r.isced == "1"][0]
        self.assertEqual(prim.completion, "full")     # "Achèvement complet"
        low = [r for r in inst.rows if r.isced == "2"][0]
        self.assertIs(low.access, True)               # "Oui, au niveau 3"
        self.assertEqual(low.orientation, "general")  # "Général"

    def test_a_second_template_family_is_read_by_its_own_profile(self):
        esp = self.by_iso.get("ESP")
        if not esp:
            self.skipTest("Spain workbook not available")
        inst, b, _ = esp
        self.assertEqual(b.profile, "uoe_scope_mapping")
        self.assertEqual(b.sheet, "Scope UOE")
        # the level is the first digit of the 3-digit ISCED-P code
        prim = [r for r in inst.rows if r.programme == "Primary education"][0]
        self.assertEqual((prim.isced, prim.isced_p), ("1", "100"))

    def test_an_age_band_is_not_a_rung_on_the_ladder(self):
        """"18-65" is adult education, not grade 13 to grade 60."""
        esp = self.by_iso.get("ESP")
        if not esp:
            self.skipTest("Spain workbook not available")
        inst = esp[0]
        off = [r for r in inst.rows if r.off_ladder]
        self.assertTrue(off)
        for r in off:
            self.assertEqual(value_of(r.grades), "-")
        gs = [s.grade for s in inst.ladder]
        self.assertLess(max(gs), 25, "the ladder ran away")

    def test_a_non_latin_script_survives(self):
        lka = self.by_iso.get("LKA")
        if not lka:
            self.skipTest("Sri Lanka workbook not available")
        inst = lka[0]
        self.assertTrue(any(any(ord(c) > 0x0D00 for c in r.programme_national)
                            for r in inst.rows))

    def test_grades_are_never_left_empty(self):
        for k, (inst, b, f) in self.by_iso.items():
            for r in inst.rows:
                self.assertIn(value_of(r.grades) is None, (False,),
                              f"{k}: {r.programme}")

    def test_nothing_blocks_on_any_of_them(self):
        for k, (inst, b, f) in self.by_iso.items():
            blocks = [x for x in run_rules(inst, b) if x.level == "BLOCK"]
            self.assertEqual(blocks, [], f"{k}: {[x.message for x in blocks]}")



# ==========================================================================
# attainment, eras, review
# ==========================================================================
class TestAttainmentVocabulary(unittest.TestCase):
    def test_a_split_level_splits(self):
        self.assertEqual(attainment_of("primary", True), "primary_complete")
        self.assertEqual(attainment_of("primary", False), "primary_incomplete")

    def test_unsettled_stays_the_bare_level_never_complete(self):
        self.assertEqual(attainment_of("primary", None), "primary")

    def test_levels_that_are_never_split(self):
        for g in ("none", "pre_primary", "other"):
            self.assertEqual(attainment_of(g, True), g)

    def test_every_produced_value_is_in_the_controlled_list(self):
        for g in GMD_EDUC_LEVELS:
            for c in (True, False, None):
                v = attainment_of(g, c)
                if v:
                    self.assertIn(v, GMD_EDUC_ATTAINMENT, f"{g}/{c}")


@needs_wb
class TestAttainmentOnRealWorkbooks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.by_iso = {}
        profs = {p["profile"]: p for p in load_profiles()}
        for f in WB:
            inst, b = read(f)
            if b.profile in profs:
                derive_run(inst, profs[b.profile])
            cls.by_iso[inst.iso3] = (inst, b, f)

    def test_every_grade_carries_an_attainment(self):
        for k, (inst, b, f) in self.by_iso.items():
            for s in inst.ladder:
                self.assertTrue(s.attainment, f"{k} grade {s.grade}")
                self.assertIn(s.attainment, GMD_EDUC_ATTAINMENT, k)

    def test_the_last_grade_of_a_level_is_complete(self):
        for k, (inst, b, f) in self.by_iso.items():
            prim = [s for s in inst.ladder if s.gmd == "primary"]
            if not prim:
                continue
            self.assertTrue(prim[-1].attainment.endswith("_complete"), k)
            self.assertTrue(prim[0].attainment.endswith("_incomplete"), k)

    def test_the_resolver_answers_with_attainment(self):
        for k, (inst, b, f) in self.by_iso.items():
            y = inst.source.reference_year
            got = resolve(inst, grade=1, year_last_in_school=y)
            self.assertEqual(got["attainment"], "primary_incomplete", k)

    def test_completion_drives_it_not_the_grade_alone(self):
        """Two programmes can share a grade and confer different attainment."""
        for k, (inst, b, f) in self.by_iso.items():
            for r in inst.rows:
                if r.completion == "full" and r.gmd not in (
                        "none", "pre_primary", "other"):
                    self.assertTrue(
                        r.attainment_on_completion.endswith("_complete"), k)
                elif r.completion in ("partial", "none") and r.gmd not in (
                        "none", "pre_primary", "other"):
                    self.assertTrue(
                        r.attainment_on_completion.endswith("_incomplete"), k)


class TestEraRegistry(unittest.TestCase):
    def _e(self, **kw):
        base = dict(iso3="ZZZ", domain="education", country="Testland",
                    file="a.xlsx", fingerprint="aaa", era="2014",
                    era_from=2014, era_to=2035, cohort_era=True,
                    signature={"grades": "1-6"})
        base.update(kw)
        return base

    def test_a_second_workbook_is_a_second_era_not_a_new_version(self):
        with tempfile.TemporaryDirectory() as d:
            REG.record(d, self._e())
            got = REG.record(d, self._e(era="2023", era_from=2023,
                                        fingerprint="bbb", file="b.xlsx"))
            self.assertEqual(got["version"], "v1.0")
            c = REG.load(d)["countries"][0]
            self.assertEqual(c["n_eras"], 2)

    def test_the_earlier_era_closes_where_the_later_one_starts(self):
        with tempfile.TemporaryDirectory() as d:
            REG.record(d, self._e())
            REG.record(d, self._e(era="2023", era_from=2023,
                                  fingerprint="bbb"))
            eras = REG.load(d)["countries"][0]["eras"]
            self.assertEqual(eras[0]["to"], 2022)
            self.assertIn("2023", eras[0]["to_inferred"])

    def test_a_changed_grade_range_is_breaking(self):
        with tempfile.TemporaryDirectory() as d:
            REG.record(d, self._e())
            got = REG.record(d, self._e(fingerprint="ccc",
                                        signature={"grades": "1-7"}))
            self.assertEqual(got["version"], "v2.0")


class TestReview(unittest.TestCase):
    def test_an_era_window_is_editable_but_a_programme_name_is_not(self):
        with tempfile.TemporaryDirectory() as d:
            RV.add_patch(d, "ZZZ_edu", dict(
                target="eras", id="1", field="from", value=1990,
                reason="an earlier mapping exists", by="A"))
            with self.assertRaises(ValueError):
                RV.add_patch(d, "ZZZ_edu", dict(
                    target="rows", id="1", field="programme", value="x",
                    reason="r", by="A"))

    def test_a_blocked_version_cannot_be_approved(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(ValueError):
                RV.approve(d, "ZZZ_edu", by="A", fingerprint="f",
                           version="v1.0", verdict="blocked",
                           counts={"BLOCK": 1})

    def test_warnings_need_acknowledging_in_writing(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(ValueError):
                RV.approve(d, "ZZZ_edu", by="A", fingerprint="f",
                           version="v1.0", verdict="awaiting_review",
                           counts={"BLOCK": 0, "WARN": 3},
                           acknowledged=[{"rule": "E-03", "text": "ok"}])


# ---------------------------------------------------------------------------
def _row(isced, name, before, at, grades="?"):
    """A programme already carrying its derived year arithmetic."""
    from gmd_edu_concordance.schema import EduRow, inferred
    r = EduRow(era=0, programme=name, isced=isced)
    r.years_before = inferred(str(before), 0.85, "test")
    r.years_at_completion = inferred(str(at), 0.85, "test")
    r.grades = inferred(grades, 0.85, "test")
    r.grade_from, r.grade_to = inferred(1, 0.85, "t"), inferred(2, 0.85, "t")
    return r


def _inst(rows):
    from gmd_edu_concordance.schema import Instruction, SourceFile
    i = Instruction(iso3="ZZZ", country="Testland",
                    source=SourceFile(file="t.xlsx", sha256="0" * 64))
    i.rows = rows
    return i


GRADE_LEVELS = ["1", "2", "3"]


class TestLadderConsistency(unittest.TestCase):
    """A starting age can be exact and still not be a rung on the ladder."""

    def test_an_in_service_entry_leaves_the_ladder(self):
        # The Central African Republic's Formation des conseillers
        # pedagogiques is ISCED 4 and admits at 30, eleven years after the
        # level below finishes. Read as a rung it produced grades 25-26.
        rows = [_row("1", "Primary", 0, 6, "1-6"),
                _row("2", "Lower secondary", 6, 10, "7-10"),
                _row("3", "In-service teacher training", 24, 26, "25-26")]
        consistency(_inst(rows), GRADE_LEVELS)
        self.assertTrue(rows[2].off_ladder)
        self.assertEqual(value_of(rows[2].grades), "-")
        self.assertIsNone(rows[2].grade_from)
        self.assertIn("rung", rows[2].grades["evidence"])
        self.assertFalse(rows[0].off_ladder)
        self.assertFalse(rows[1].off_ladder)

    def test_a_genuine_extra_rung_stays(self):
        # Nigeria's IJMB A-level course starts the year senior secondary
        # ends. That is a continuation, not a second-career entry.
        rows = [_row("1", "Primary", 0, 6, "1-6"),
                _row("2", "Junior secondary", 6, 9, "7-9"),
                _row("3", "Senior secondary", 9, 12, "10-12"),
                _row("3", "IJMB A-level course", 12, 14, "13-14")]
        consistency(_inst(rows), GRADE_LEVELS)
        self.assertFalse(rows[3].off_ladder)
        self.assertEqual(value_of(rows[3].grades), "13-14")

    def test_a_dropped_row_does_not_raise_the_level_above(self):
        # Argentina: Adults primary education finishes at year 12. If that
        # counted as the top of ISCED 1, adult lower secondary entering at
        # year 12 looked perfectly normal and kept grades 13-14.
        rows = [_row("1", "Primary education", 0, 6, "1-6"),
                _row("1", "Adults primary education", 9, 12, "10-12"),
                _row("2", "Basic cycle", 6, 9, "7-9"),
                _row("2", "Basic cycle, adults", 12, 14, "13-14")]
        consistency(_inst(rows), GRADE_LEVELS)
        self.assertTrue(rows[1].off_ladder)
        self.assertTrue(rows[3].off_ladder,
                        "a row ruled off the ladder must not go on to raise "
                        "the baseline for the level above it")

    def test_a_named_adult_route_is_never_a_rung(self):
        # Numerically identical to the IJMB course above; only the name
        # separates them, and the workbook gives the name.
        rows = [_row("1", "Primary", 0, 6, "1-6"),
                _row("2", "Basic cycle", 6, 9, "7-9"),
                _row("3", "Oriented cycle", 9, 12, "10-12"),
                _row("3", "Oriented cycle. Adults secondary education",
                     12, 14, "13-14")]
        consistency(_inst(rows), GRADE_LEVELS)
        self.assertTrue(rows[3].off_ladder)
        self.assertIn("adult", rows[3].grades["evidence"].lower())

    def test_the_marker_is_read_in_the_country_language(self):
        from gmd_edu_concordance.schema import EduRow, inferred
        r = _row("2", "", 6, 9, "7-9")
        r.programme_national = "Educación de jóvenes y adultos"
        rows = [_row("1", "Primary", 0, 6, "1-6"), r]
        consistency(_inst(rows), GRADE_LEVELS)
        self.assertTrue(r.off_ladder)

    def test_tertiary_is_left_alone(self):
        rows = [_row("1", "Primary", 0, 6, "1-6"),
                _row("6", "Bachelor", 12, 16, "-")]
        consistency(_inst(rows), GRADE_LEVELS)
        self.assertFalse(rows[1].off_ladder)


@needs_wb
class TestSheetSelection(unittest.TestCase):
    """More than one sheet, and the mapping table is not always the first."""

    @staticmethod
    def _with_decoys(src, dst, header_row):
        """Three sheets ahead of the real one, each wrong in its own way.

        The empty annex is the dangerous decoy: it carries the *same header
        row* as the mapping table, so it binds every required column and the
        first-sheet-that-binds rule would read it and emit a valid, empty
        instruction.
        """
        import openpyxl
        wb = openpyxl.load_workbook(src)
        real = wb.sheetnames[0]
        src_ws = wb[real]

        ins = wb.create_sheet("Instructions", 0)
        ins["A1"] = ("Please read these instructions carefully before "
                     "completing the mapping. Each programme in the national "
                     "education system should be entered on one row.")
        ins["A3"] = ("The ISCED level is assigned according to the criteria "
                     "set out in the ISCED 2011 operational manual.")

        lists = wb.create_sheet("Dropdown_lists", 1)
        for i, v in enumerate(["ISCED level", "0", "1", "2", "3"], start=1):
            lists.cell(row=i, column=1, value=v)

        annex = wb.create_sheet("Annex - blank", 2)
        for c in range(1, (src_ws.max_column or 1) + 1):
            annex.cell(row=1, column=c,
                       value=src_ws.cell(row=header_row, column=c).value)
        annex["A2"] = "n/a"

        wb.save(dst)
        return real

    def test_the_mapping_table_wins_over_prose_and_lists(self):
        with tempfile.TemporaryDirectory() as d:
            dst = os.path.join(d, os.path.basename(WB[0]))
            plain, pb = read(WB[0])
            real = self._with_decoys(WB[0], dst, pb.header_row)
            inst, b = read(dst)
            self.assertEqual(b.sheet, real)
            self.assertEqual(len(inst.rows), len(plain.rows))
            names = [c["sheet"] for c in (b.candidates_considered or [])]
            self.assertIn("Annex - blank", names,
                          "the blank annex binds every required column and "
                          "must be scored, not skipped")
            self.assertGreater(len(names), 1)

    def test_the_sheet_can_be_forced(self):
        with tempfile.TemporaryDirectory() as d:
            dst = os.path.join(d, os.path.basename(WB[0]))
            _plain, pb = read(WB[0])
            real = self._with_decoys(WB[0], dst, pb.header_row)
            inst, b = read(dst, sheet="Annex - blank")
            self.assertEqual(b.sheet, "Annex - blank")
            inst, b = read(dst, sheet=real)
            self.assertEqual(b.sheet, real)

    def test_prose_scores_below_a_table(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            dst = os.path.join(d, os.path.basename(WB[0]))
            _plain, pb = read(WB[0])
            self._with_decoys(WB[0], dst, pb.header_row)
            wb = openpyxl.load_workbook(dst)
            prof = load_profile()
            best = {}
            for name in wb.sheetnames:
                sc = score_sheet(wb[name], prof, name)
                best[name] = sc["score"] if sc else -1
            top = max(best.values())
            self.assertLess(best["Instructions"], top)
            self.assertLess(best["Dropdown_lists"], top)
            self.assertGreater(top, 0)


if __name__ == "__main__":
    if not WB:
        print("note: no ISCED workbook found -- workbook tests will be "
              "skipped.\n      set GMD_ISCED_DIR=/path/to/files to run them.\n")
    unittest.main(verbosity=2)
