"""
gmd_wash_concordance.reader
===========================
JMP country workbook -> Instruction.

The reader decides nothing a person would want to decide, and it never drops a
row silently: anything it cannot place lands in ``instruction.unmapped`` with
the source row and a reason.

Layout, for the record
----------------------
Each ``Water Data`` / ``Sanitation Data`` sheet is a strip of fixed-width
blocks, one per data source, repeating across the sheet::

    col B          C                    D              E      F      G
    Definitions    Facility type ...    (category)     Urban  Rural  Total
    ^ the country's own denomination    ^ the JMP master tree

Row 2 of a block carries the source code (``IND_2016_DHS``), row 3 its type
(``Survey with microdata``) and, two columns over, its name.  Somewhere below
sits a cell reading *Original denomination*; everything under it, for the
length of the master tree, is the concordance this source supplies.
"""
from __future__ import annotations

import os
from dataclasses import asdict
import re
import unicodedata
from typing import Dict, List, Optional, Tuple

import yaml
from openpyxl import load_workbook

from . import benchmark as BM
from . import countries as CO
from . import master as M
from . import translate as TR
from .schema import (
    DataSource, Extract, Instruction, Section, SourceConcordance, SourceFile,
    Unmapped, WashRow,
)

PROFILE_DIR = os.path.join(os.path.dirname(__file__), "data")
DEFAULT_PROFILE = "jmp_country_file_2025.yaml"


# --------------------------------------------------------------------------
def _norm(s) -> str:
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().split())


def _txt(ws, r, c) -> str:
    v = ws.cell(r, c).value
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _num(ws, r, c):
    v = ws.cell(r, c).value
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _en(ws, fws, r, c, tr) -> str:
    """The English phrase in a cell, resolved exactly where possible.

    The formula names the translation-table row, which identifies the phrase;
    the rendered string is only a fallback, because translations are
    many-to-one and a string map has to guess between the collisions.
    """
    raw = _txt(ws, r, c)
    if tr is None or not raw:
        return raw
    fml = fws.cell(r, c).value if fws is not None else None
    return tr.en_of(raw, fml)


def slug(s: str) -> str:
    s = _norm(s)
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    return s or "x"


def load_profile(path: Optional[str] = None) -> dict:
    with open(path or os.path.join(PROFILE_DIR, DEFAULT_PROFILE),
              encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def _pick_sheet(wb, names: List[str]) -> Optional[str]:
    have = {_norm(t): t for t in wb.sheetnames}
    for n in names:
        if _norm(n) in have:
            return have[_norm(n)]
    return None


# --------------------------------------------------------------------------
class Binding:
    """What bound where.  ``inspect`` prints this and stops."""

    def __init__(self, profile_name: str):
        self.profile = profile_name
        self.sheets: Dict[str, str] = {}
        self.tree_header_row: Dict[str, Optional[int]] = {}
        self.tree_alignment: Dict[str, List[dict]] = {}
        self.blocks: Dict[str, int] = {}
        self.selected: Dict[str, int] = {}
        self.sections: Dict[str, List[dict]] = {}
        self.translation: Dict[str, Any] = {}
        self.identity: Dict[str, str] = {}
        self.problems: List[str] = []

    def as_dict(self):
        return {"profile": self.profile, "identity": self.identity,
                "sheets": self.sheets, "tree_header_row": self.tree_header_row,
                "blocks": self.blocks, "selected": self.selected,
                "sections": self.sections,
                "translation": self.translation,
                "tree_alignment": self.tree_alignment,
                "problems": self.problems}


# --------------------------------------------------------------------------
def _identity_language(wb, prof) -> str:
    sheet = _pick_sheet(wb, prof["identity"]["sheet"])
    if not sheet:
        return ""
    try:
        return str(wb[sheet][prof["identity"]["language_cell"]].value or "").strip()
    except Exception:                                      # noqa: BLE001
        return ""


def _identity(wb, path, prof) -> Tuple[str, str, str]:
    """(iso3, country, language) -- from the Introduction sheet, falling back
    to the filename convention."""
    iso3 = country = lang = ""
    sheet = _pick_sheet(wb, prof["identity"]["sheet"])
    if sheet:
        ws = wb[sheet]
        try:
            country = str(ws[prof["identity"]["country_name_cell"]].value or "").strip()
            iso3 = str(ws[prof["identity"]["iso3_cell"]].value or "").strip().upper()
            lang = str(ws[prof["identity"]["language_cell"]].value or "").strip()
        except Exception:
            pass
    if not iso3 or len(iso3) != 3 or not CO.english_name(iso3):
        # The Introduction sheet gives the code directly in every file seen so
        # far, but a translated or hand-edited file may not, and the country
        # name it carries is in the file's own language.
        got, name = CO.resolve(country, str(path))
        iso3 = got or iso3
        country = country or name
    return iso3, country, lang


def _vintage_from_name(path) -> str:
    m = re.search(r"JMP_(\d{4})_", os.path.basename(str(path)))
    return m.group(1) if m else ""


def _parse_code(code: str) -> Tuple[Optional[int], str]:
    """``IND_2016_DHS`` -> (2016, 'DHS')."""
    parts = code.split("_")
    year = None
    series = ""
    for p in parts[1:]:
        if year is None and re.fullmatch(r"(19|20)\d{2}", p):
            year = int(p)
        elif year is not None:
            series = p if not series else series + "_" + p
    return year, series


def _find_tree_header(ws, col: int, anchors: List[str], limit: int = 400,
                      tr: Optional[TR.Dictionary] = None,
                      fws=None) -> Optional[int]:
    """Find an anchor row, in any language.

    Each candidate cell is resolved back to English through the workbook's own
    translation table before it is compared, so the anchors in the profile stay
    English and a Russian file still binds.
    """
    want = [_norm(a) for a in anchors]
    for r in range(1, min(ws.max_row or 1, limit) + 1):
        raw = _txt(ws, r, col)
        if not raw:
            continue
        cands = {raw}
        if tr is not None:
            cands.add(tr.en(raw))
            cands.add(_en(ws, fws, r, col, tr))
        for cand in cands:
            v = _norm(cand)
            if v and any(v.startswith(w) for w in want):
                return r
    return None


def _read_sections(ws, c_def: int, c_class: int, c_cat: int,
                   first_row: int, stop_row: int, values=False,
                   c_vals=(), tr: Optional[TR.Dictionary] = None,
                   fws=None) -> List[Section]:
    """The blocks stacked above the concordance table.

    A value in the classification column opens a section; the values in the
    category column below it are its items.  *Facility type estimates* is the
    one that matters -- it is the last link in the chain a reviewer follows,
    classification -> subgroup -> facility type -- and reading it from the
    workbook rather than pinning it keeps the country file's own wording, in
    the country file's own language.
    """
    out: List[Section] = []
    cur: Optional[Section] = None
    for r in range(first_row, stop_row):
        head = _txt(ws, r, c_class)
        item = _txt(ws, r, c_cat)
        if head and not item:
            en = _en(ws, fws, r, c_class, tr) if tr is not None else head
            cur = Section(name=en, name_local=head, row=r)
            out.append(cur)
            continue
        if item and cur is not None:
            en = _en(ws, fws, r, c_cat, tr) if tr is not None else item
            rec: Dict[str, Any] = {"label": en, "label_local": item, "row": r}
            if values:
                for name, col in zip(("urban", "rural", "total"), c_vals):
                    v = _txt(ws, r, col)
                    if v:
                        rec[name] = v
            cur.items.append(rec)
    return out


def _tree_labels(ws, c_class: int, c_cat: int, first: int, n: int,
                 tr: Optional[TR.Dictionary] = None, fws=None) -> List[str]:
    out = []
    for i in range(n):
        r = first + i
        col = c_class if _txt(ws, r, c_class) else c_cat
        out.append(_en(ws, fws, r, col, tr) if tr is not None
                   else (_txt(ws, r, c_class) or _txt(ws, r, c_cat)))
    return out


# --------------------------------------------------------------------------
def _chain(node) -> Dict[str, Any]:
    """The chain a reviewer follows: classification -> subgroup -> facility
    type.  ``path`` is ``"Ground water > Protected well"``; the classification
    is the group at the top of it and the subgroup is everything under."""
    parts = [p.strip() for p in node.path.split(">")]
    return {
        "classification": parts[0],
        "subgroup": " > ".join(parts[1:]),
        "depth": node.level,
        "rolls_up_to": list(node.rolls_up_to),
    }


def _read_domain(ws, prof, tree: M.MasterTree, binding: Binding, domain: str,
                 want_types, unmapped: List[Unmapped],
                 tr: Optional[TR.Dictionary] = None, fws=None
                 ) -> Tuple[List[DataSource], Dict[str, List[WashRow]],
                            Dict[str, List[dict]], List[Section], Dict[str, dict]]:
    blk = prof["block"]
    first_col, width, off = blk["first_column"], blk["width"], blk["offsets"]
    code_row, type_row = blk["source_code_row"], blk["source_type_row"]
    name_off = blk["source_name_offset"]
    anchors = prof["anchor"]["concordance_header"]
    used_anchor = prof["anchor"].get("data_used_header", ["data used for estimates"])

    max_col = ws.max_column or 0
    n_blocks = max(0, (max_col - first_col + 1) // width)
    binding.blocks[domain] = 0
    binding.selected[domain] = 0

    sources: List[DataSource] = []
    rows_by_source: Dict[str, List[WashRow]] = {}
    used_by_source: Dict[str, List[dict]] = {}
    stats_by_source: Dict[str, dict] = {}
    sections: List[Section] = []
    header_row = None

    for b in range(n_blocks):
        c0 = first_col + b * width
        code = _txt(ws, code_row, c0)
        if not code:
            continue
        binding.blocks[domain] += 1
        stype = _txt(ws, type_row, c0)
        name = _txt(ws, type_row, c0 + name_off)
        stype_en = tr.en(stype) if tr is not None else stype
        if want_types is not None and _norm(stype_en) not in want_types:
            continue

        c_def = c0 + off["definitions"]
        c_class = c0 + off["classification"]
        c_cat = c0 + off["category"]
        c_vals = (c0 + off["urban"], c0 + off["rural"], c0 + off["total"])
        hdr = _find_tree_header(ws, c_def, anchors, tr=tr, fws=fws)
        if hdr is None:
            unmapped.append(Unmapped(
                where=f"{domain}.block", source=code, source_row=None, value=code,
                reason="no 'Original denomination' anchor in this block"))
            continue

        if header_row is None:
            header_row = hdr
            binding.tree_header_row[domain] = hdr
            labels = _tree_labels(ws, c_class, c_cat, hdr + 1, len(tree), tr, fws)
            binding.tree_alignment[domain] = M.check_alignment(tree, labels)
            # stop above the free-text Notes row, or the note becomes a
            # section with no items
            notes_at = _find_tree_header(ws, c_def, ["notes"], limit=hdr,
                                         tr=tr, fws=fws)
            sections = _read_sections(ws, c_def, c_class, c_cat,
                                      type_row + 1, notes_at or hdr, tr=tr,
                                      fws=fws)
            binding.sections[domain] = [
                {"name": x.name, "row": x.row, "items": len(x.items)}
                for x in sections]

        # The aggregate rows sit at the same offsets in every block, but the
        # rows are read per block so a block whose Notes run longer is still
        # read from its own anchors rather than the first block's.
        notes_here = _find_tree_header(ws, c_def, ["notes"], limit=hdr, tr=tr,
                                       fws=fws)
        block_sections = _read_sections(ws, c_def, c_class, c_cat,
                                        type_row + 1, notes_here or hdr,
                                        tr=tr, fws=fws)

        # the Yes/No flags saying which estimates JMP took from this source
        used: List[dict] = []
        u_at = _find_tree_header(ws, c_class, used_anchor, limit=hdr, tr=tr,
                                 fws=fws)
        if u_at:
            for sec in _read_sections(ws, c_def, c_class, c_cat, u_at, hdr,
                                      values=True, c_vals=c_vals, tr=tr,
                                      fws=fws):
                used.extend(sec.items)

        year, series = _parse_code(code)
        src = DataSource(code=code, year=year, series=series, type=stype_en,
                         name=name, column=c0)
        rows: List[WashRow] = []
        for i, node in enumerate(tree.nodes):
            r = hdr + 1 + i
            denom = _txt(ws, r, c_def)
            if not denom:
                continue
            if node.structural:
                unmapped.append(Unmapped(
                    where=f"{domain}.rows", source=code, source_row=r, value=denom,
                    reason=f"national category sits on the structural row "
                           f"'{node.label}', which is not a JMP class"))
                continue
            ch = _chain(node)
            rows.append(WashRow(
                code=slug(denom), label=denom,
                jmp=node.path, jmp_id=node.id,
                gmd=node.gmd, spans=list(node.spans),
                classification=ch["classification"], subgroup=ch["subgroup"],
                depth=ch["depth"], rolls_up_to=ch["rolls_up_to"],
                improved=node.improved, shared=node.shared,
                observed_in=[code],
                first_seen=year, last_seen=year,
                source_row=r,
                jmp_label_local=_txt(ws, r, c_class) or _txt(ws, r, c_cat),
                language=(tr.language if tr is not None else ""),
            ))
        src.n_rows = len(rows)
        sources.append(src)
        rows_by_source[code] = rows
        used_by_source[code] = used
        stats_by_source[code] = BM.for_source(
            ws, tree, hdr, c_def, c_class, c_cat, c_vals,
            block_sections if block_sections else sections, tr, fws)
        binding.selected[domain] += 1

    if header_row is None and binding.blocks.get(domain):
        binding.tree_header_row[domain] = None
        binding.problems.append(
            f"no block on the {domain} sheet carried a concordance header. "
            f"The anchor is matched in English through the workbook's own "
            f"translation table; if that table could not be read, the labels "
            f"cannot be resolved.")
    return sources, rows_by_source, used_by_source, sections, stats_by_source


def _merge(rows_by_source: Dict[str, List[WashRow]],
           order: List[str]) -> List[WashRow]:
    """One row per (national label, JMP node), carrying the sources that used
    it and the years it spans.  That list *is* the vintage trail."""
    merged: Dict[Tuple[str, str], WashRow] = {}
    for code in order:
        for r in rows_by_source.get(code, []):
            key = (_norm(r.label), r.jmp_id)
            cur = merged.get(key)
            if cur is None:
                merged[key] = WashRow(**{**r.__dict__,
                                         "spans": list(r.spans),
                                         "observed_in": list(r.observed_in)})
            else:
                cur.observed_in.extend(r.observed_in)
                ys = [y for y in (cur.first_seen, cur.last_seen, r.first_seen,
                                  r.last_seen) if y is not None]
                cur.first_seen = min(ys) if ys else None
                cur.last_seen = max(ys) if ys else None
    out = list(merged.values())
    out.sort(key=lambda r: (r.jmp_id, _norm(r.label)))
    return out


# --------------------------------------------------------------------------
def read(path, iso3: Optional[str] = None, country: Optional[str] = None,
         version: str = "v1.0", sources: str = "microdata",
         domains: str = "both",
         profile_path: Optional[str] = None) -> Tuple[Extract, Binding]:
    """Read one JMP country workbook into **two** instructions.

    ``sources`` selects which blocks to extract:
      ``microdata`` (default) | ``survey`` | ``census`` | ``admin`` | ``all``
      or a comma-separated combination.

    ``domains`` selects which to build: ``both`` (default) | ``water`` |
    ``sanitation``.  They are separate instructions, separate files and
    separate approvals.
    """
    prof = load_profile(profile_path)
    binding = Binding(prof["profile"])
    wb = load_workbook(path, data_only=True, read_only=False)

    tr = TR.load(wb, _identity_language(wb, prof))
    # A workbook whose labels are already English needs no second read; one in
    # any other language does, because only the formula identifies the phrase.
    wbf = None
    if tr.available and not tr.is_english:
        wbf = load_workbook(path, data_only=False, read_only=False)
    fiso, fcountry, lang = _identity(wb, path, prof)
    iso3 = (iso3 or fiso or "").upper()
    country = country or fcountry
    src_file = SourceFile.of(path, prof["profile"], language=lang,
                             vintage=_vintage_from_name(path))
    binding.identity = {"iso3": iso3, "country": country,
                        "language": lang, "vintage": src_file.vintage,
                        "country_en": CO.english_name(iso3)}
    binding.translation = {
        "sheet": tr.sheet, "language": tr.language or lang,
        "languages": tr.languages, "phrases": len(tr),
        "available": tr.available,
        "resolved_by": ("formula index" if wbf is not None
                        else "already English"),
        "ambiguous_strings": tr.ambiguous,
    }
    if not tr.available and _norm(lang) not in ("", "english"):
        binding.problems.append(
            f"the workbook is in {lang} and its translation table could not be "
            f"read, so its labels cannot be resolved to English")
    if not iso3:
        binding.problems.append(
            "country could not be determined from the Introduction sheet or "
            "the filename -- pass --iso3")

    want = None
    if sources != "all":
        want = set()
        for key in [x.strip() for x in sources.split(",") if x.strip()]:
            if key not in prof["source_types"]:
                raise ValueError(
                    f"unknown source selector {key!r}; expected one of "
                    f"{sorted(prof['source_types'])} or 'all'")
            want |= {_norm(v) for v in prof["source_types"][key]}
    selection = {"sources": sources, "types": sorted(want) if want else "all"}

    wanted_domains = (["water", "sanitation"] if domains == "both"
                      else [d.strip() for d in domains.split(",") if d.strip()])
    for d in wanted_domains:
        if d not in prof["domains"]:
            raise ValueError(f"unknown domain {d!r}; expected water, "
                             f"sanitation or both")

    ex = Extract(iso3=iso3, country=country, language=lang)
    ladder_bm = BM.ladder_estimates(wb)

    for domain in wanted_domains:
        cfg = prof["domains"][domain]
        sheet = _pick_sheet(wb, cfg["sheet"])
        binding.sheets[domain] = sheet or ""
        if sheet is None:
            binding.problems.append(
                f"no sheet matched {cfg['sheet']} for {domain}")
            continue
        tree = M.load(cfg["master"])
        inst = Instruction(
            domain=domain, iso3=iso3, country=country, version=version,
            frame=prof["frame"], language=lang, source=src_file,
            selection=dict(selection),
            translation=dict(binding.translation),
        )
        srcs, rows_by, used_by, sections, stats_by = _read_domain(
            wb[sheet], prof, tree, binding, domain, want, inst.unmapped, tr,
            wbf[sheet] if wbf is not None else None)

        srcs.sort(key=lambda s: (s.year or 0, s.code))
        inst.sources = srcs
        inst.rows = _merge(rows_by, [s.code for s in srcs])
        inst.sections = sections
        inst.per_source = [
            SourceConcordance(source=s, rows=rows_by.get(s.code, []),
                              data_used=used_by.get(s.code, []))
            for s in srcs
        ]
        inst.facility_types = _facility_types(tree, sections)
        inst.classifications = _classifications(tree, inst.rows)

        inst.benchmark = {
            "country": {"iso3": iso3, "name": country},
            "domain": domain,
            "frame": prof["frame"],
            "language": lang,
            "file": asdict(src_file),
            "note": "published by JMP; kept beside the instruction and never "
                    "inside it -- a country instruction may not carry "
                    "survey-level data",
            "ladder": ladder_bm,
            "per_source": [
                {"source": asdict(s), **stats_by.get(s.code, {})}
                for s in srcs
            ],
        }
        inst.benchmark["summary"] = BM.summarise(inst.benchmark)

        years = [s.year for s in srcs if s.year]
        inst.vintage = {
            "from": min(years) if years else None,
            "to": max(years) if years else None,
            "n_sources": len(srcs),
            "file_release": src_file.vintage,
        }
        setattr(ex, domain, inst)

    return ex, binding


def _facility_types(tree: M.MasterTree, sections: List[Section]
                    ) -> List[Dict[str, Any]]:
    """The workbook's own *Facility type estimates* rows, each with the JMP
    classes that feed it.  This is the last link in the chain."""
    local = {}
    for sec in sections:
        if _norm(sec.name).startswith("facility type"):
            for it in sec.items:
                local[_norm(it["label"])] = it.get("label_local") or it["label"]
    out = []
    for name in tree.facility_types:
        feeders = [n.path for n in tree.nodes if name in n.rolls_up_to]
        out.append({
            "name": name,
            "name_local": local.get(_norm(name), ""),
            "in_workbook": _norm(name) in local,
            "n_classes": len(feeders),
            "classes": feeders,
        })
    return out


def _classifications(tree: M.MasterTree, rows: List[WashRow]
                     ) -> List[Dict[str, Any]]:
    """The classification groups, their subgroups, and which of them the
    country actually used.  The tree is the JMP master's; the counts are this
    country's."""
    used = {}
    for r in rows:
        used.setdefault(r.jmp_id, []).append(r)
    out = []
    group = None
    for n in tree.nodes:
        if n.structural:
            continue
        if n.level == 0:
            group = {"classification": n.label, "jmp_id": n.id,
                     "improved": n.improved, "gmd": n.gmd,
                     "spans": list(n.spans),
                     "rolls_up_to": list(n.rolls_up_to),
                     "n_national_categories": len(used.get(n.id, [])),
                     "subgroups": []}
            out.append(group)
            continue
        if group is None:
            continue
        group["subgroups"].append({
            "label": n.label, "jmp_id": n.id, "depth": n.level,
            "improved": n.improved, "gmd": n.gmd, "spans": list(n.spans),
            "shared": n.shared, "rolls_up_to": list(n.rolls_up_to),
            "national_categories": [r.label for r in used.get(n.id, [])],
        })
    for g in out:
        g["n_national_categories"] += sum(
            len(sg["national_categories"]) for sg in g["subgroups"])
    return out
