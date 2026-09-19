from __future__ import annotations

import argparse
import json
import re
from hashlib import sha256
from pathlib import Path
from typing import Any

import yaml

from .emit import emit_benchmark, emit_parameter_draft
from ..hashing import hash_file
from .legacy import ROOT, load_edu_module, load_wash_module
from .transform import (
    to_benchmark_payload,
    transform_education_rows,
    transform_wash_rows,
)
from schema.country_parameters import CountryParameterFile
from schema.frontmatter import load_markdown
from schema.parameter import ParameterDefinition

SOURCE_ROOT = ROOT / "extraction" / "10_source" / "country-parameters-inputs"
DRAFT_ROOT = ROOT / "extraction" / "20_drafts" / "runs" / "country-parameters"
CONTRACT_ROOT = DRAFT_ROOT / "contracts"
STATE_PATH = DRAFT_ROOT / ".incremental_state.json"

ISO3_PATTERN = re.compile(r"^[A-Z]{3}$")


def _is_valid_iso3(value: str) -> bool:
    return bool(ISO3_PATTERN.fullmatch(value or ""))


def _delete_if_exists(path: Path) -> bool:
    if path.exists():
        path.unlink()
        return True
    return False


def _cleanup_stale_root_level_yaml() -> list[str]:
    cleaned: list[str] = []
    for path in sorted(DRAFT_ROOT.glob("*.yaml")):
        path.unlink()
        cleaned.append(str(path.relative_to(ROOT)))
    return cleaned


def _load_incremental_state() -> dict[str, Any]:
    if not STATE_PATH.exists():
        return {"version": 1, "artifacts": {}}
    try:
        data = json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {"version": 1, "artifacts": {}}
    if not isinstance(data, dict):
        return {"version": 1, "artifacts": {}}
    if not isinstance(data.get("artifacts"), dict):
        data["artifacts"] = {}
    data.setdefault("version", 1)
    return data


def _save_incremental_state(state: dict[str, Any]) -> None:
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(json.dumps(state, indent=2, ensure_ascii=False), encoding="utf-8")


def _signature(*parts: str) -> str:
    joined = "\n".join(parts).encode("utf-8")
    return sha256(joined).hexdigest()


def _file_sha256(path: Path) -> str:
    if not path.exists():
        return "missing"
    return hash_file(path)


def _extractor_signature() -> str:
    files = [
        ROOT / "extraction_pipeline" / "country_inputs" / "cli.py",
        ROOT / "extraction_pipeline" / "country_inputs" / "transform.py",
        ROOT / "extraction_pipeline" / "country_inputs" / "emit.py",
        ROOT / "schema" / "country_parameters.py",
        ROOT / "schema" / "parameter.py",
    ]
    return _signature(*[_file_sha256(path) for path in files])


def _parameter_schema_signature(parameter_id: str) -> str:
    contract_path = ROOT / "knowledge" / "parameters" / f"{parameter_id}.md"
    return _signature(parameter_id, _file_sha256(contract_path))


def _should_emit(
    *,
    state: dict[str, Any],
    key: str,
    signature: str,
    output_path: Path,
    force: bool,
) -> bool:
    if force:
        return True
    artifacts = state.get("artifacts", {})
    entry = artifacts.get(key)
    if not isinstance(entry, dict):
        return True
    if not output_path.exists():
        return True
    return entry.get("signature") != signature


def _record_artifact_state(
    state: dict[str, Any],
    key: str,
    signature: str,
    output_path: Path,
) -> None:
    artifacts = state.setdefault("artifacts", {})
    artifacts[key] = {
        "signature": signature,
        "output": str(output_path.relative_to(ROOT)),
    }


def _drop_artifact_state(state: dict[str, Any], key: str) -> None:
    artifacts = state.get("artifacts", {})
    if isinstance(artifacts, dict):
        artifacts.pop(key, None)


def _find_xlsx(folder: Path) -> list[Path]:
    return sorted(path for path in folder.glob("*.xlsx") if not path.name.startswith("~$"))


def _to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _infer_subnat_level(byvar: str) -> int:
    if byvar == "subnatid":
        return 1
    match = re.fullmatch(r"subnatid(\d+)", byvar)
    if match:
        return int(match.group(1))
    return 0


def _country_name_for_iso3(iso3: str) -> str:
    country_file = ROOT / "country-parameters" / "countries" / iso3 / "parameters.md"
    if not country_file.exists():
        return iso3
    try:
        data, _ = load_markdown(country_file)
    except Exception:  # noqa: BLE001
        return iso3
    if isinstance(data, dict):
        name = data.get("country_name")
        if isinstance(name, str) and name.strip():
            return name.strip()
    return iso3


def _to_int_year(value: Any) -> int | None:
    text = _to_text(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _geo_boundary_signature(row: dict[str, Any]) -> tuple[str, ...]:
    return (
        row["survey_variables"],
        row["gmd_subnatid1"],
        row["gmd_subnatid2"],
        row["gmd_subnatid3"],
        row["gmd_subnatid4"],
        row["gmd_subnatidsurvey"],
        row["geo_source"],
        row["geo_level"],
        row["geo_idvar"],
        row["geo_id"],
    )


def _active_geo_code(row: dict[str, Any]) -> str:
    for key in (
        "gmd_subnatid1",
        "gmd_subnatid2",
        "gmd_subnatid3",
        "gmd_subnatid4",
        "gmd_subnatidsurvey",
    ):
        value = _to_text(row.get(key))
        if value:
            return value
    return ""


def _variant_tokens(value: Any) -> set[str]:
    text = _to_text(value)
    if not text:
        return set()
    return {part.strip() for part in text.split("|") if part.strip()}


def _dedupe_geo_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    # Uniqueness is anchored on geo_code identity and structural scope, not survey_label text.
    # This allows multiple survey_label spellings/variants for the same geo_code without
    # emitting duplicate rows.
    seen: dict[tuple[str, ...], dict[str, Any]] = {}
    order: list[tuple[str, ...]] = []
    for row in rows:
        geo_code = _active_geo_code(row)
        key = (
            row.get("geo_year", ""),
            row.get("geo_source", ""),
            row.get("geo_level", ""),
            row.get("geo_idvar", ""),
            row.get("geo_id", ""),
            row.get("geo_nvar", ""),
            geo_code,
        )
        existing = seen.get(key)
        if existing is None:
            seen[key] = row
            order.append(key)
            continue

        # Merge survey label/variable variants and structural fields by non-empty preference.
        labels = _variant_tokens(existing.get("survey_labels")) | _variant_tokens(
            row.get("survey_labels")
        )
        variables = _variant_tokens(existing.get("survey_variables")) | _variant_tokens(
            row.get("survey_variables")
        )
        existing["survey_labels"] = " | ".join(sorted(labels)) if labels else ""
        existing["survey_variables"] = " | ".join(sorted(variables)) if variables else ""

        for gmd_key in (
            "gmd_subnatid1",
            "gmd_subnatid2",
            "gmd_subnatid3",
            "gmd_subnatid4",
            "gmd_subnatidsurvey",
        ):
            if not _to_text(existing.get(gmd_key)) and _to_text(row.get(gmd_key)):
                existing[gmd_key] = row[gmd_key]

        for rep_key in ("is_rep_subnat1", "is_rep_subnat2", "is_rep_subnat3", "is_rep_subnat4"):
            existing[rep_key] = bool(existing.get(rep_key)) or bool(row.get(rep_key))

        try:
            existing_level = int(existing.get("representative_level", 0))
        except (TypeError, ValueError):
            existing_level = 0
        try:
            row_level = int(row.get("representative_level", 0))
        except (TypeError, ValueError):
            row_level = 0
        existing["representative_level"] = max(existing_level, row_level)

        src_existing = existing.get("source_row")
        src_new = row.get("source_row")
        if isinstance(src_existing, int) and isinstance(src_new, int) and src_new < src_existing:
            # Keep earliest source row for traceability while preserving merged variants.
            existing["source_row"] = src_new

    return [seen[key] for key in order]


def _build_geo_parameter_records(rows: list[dict[str, Any]], source: str) -> list[dict[str, Any]]:
    by_year: dict[int, list[dict[str, Any]]] = {}
    unknown_year_rows: list[dict[str, Any]] = []
    for row in rows:
        year = _to_int_year(row.get("geo_year"))
        if year is None:
            unknown_year_rows.append(row)
            continue
        by_year.setdefault(year, []).append(row)

    if not by_year:
        return [
            {
                "parameter_id": "PARAM-GEO-GMD-CROSSWALK",
                "effective_from": None,
                "effective_to": None,
                "selectors": None,
                "value": rows,
                "provenance": {
                    "source": source,
                    "verified_on": None,
                    "human_reviewed": False,
                    "reviewer": None,
                },
            }
        ]

    years = sorted(by_year.keys())
    signatures_by_year: dict[int, tuple[tuple[str, ...], ...]] = {}
    for year in years:
        sigs = sorted({_geo_boundary_signature(row) for row in by_year[year]})
        signatures_by_year[year] = tuple(sigs)

    segments: list[dict[str, int]] = []
    start_year = years[0]
    prev_year = years[0]
    prev_sig = signatures_by_year[prev_year]
    for year in years[1:]:
        sig = signatures_by_year[year]
        if sig != prev_sig:
            segments.append(
                {
                    "start": start_year,
                    "end": prev_year,
                    "representative": start_year,
                }
            )
            start_year = year
        prev_year = year
        prev_sig = sig

    segments.append(
        {
            "start": start_year,
            "end": prev_year,
            "representative": start_year,
        }
    )

    records: list[dict[str, Any]] = []
    for idx, segment in enumerate(segments):
        record_rows = by_year[segment["representative"]]
        effective_to = None if idx == len(segments) - 1 else segment["end"]
        records.append(
            {
                "parameter_id": "PARAM-GEO-GMD-CROSSWALK",
                "effective_from": segment["start"],
                "effective_to": effective_to,
                "selectors": None,
                "value": record_rows,
                "provenance": {
                    "source": source,
                    "verified_on": None,
                    "human_reviewed": False,
                    "reviewer": None,
                },
            }
        )

    if unknown_year_rows:
        records.append(
            {
                "parameter_id": "PARAM-GEO-GMD-CROSSWALK",
                "effective_from": None,
                "effective_to": None,
                "selectors": {"geo_year": "unknown"},
                "value": unknown_year_rows,
                "provenance": {
                    "source": source,
                    "verified_on": None,
                    "human_reviewed": False,
                    "reviewer": None,
                },
            }
        )

    return records


def _emit_geo_parameter_draft(iso3: str, country_name: str, rows: list[dict[str, Any]], source: str) -> Path:
    out_dir = DRAFT_ROOT / iso3
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "PARAM-GEO-GMD-CROSSWALK.yaml"

    payload = {
        "country_id": f"CTY-{iso3}",
        "country_name": country_name,
        "iso3": iso3,
        "schema_version": "0.2",
        "status": "draft",
        "parameters": _build_geo_parameter_records(rows, source),
    }

    text = yaml.safe_dump(payload, sort_keys=False, allow_unicode=False)
    out_path.write_text(text, encoding="utf-8")
    return out_path


def extract_geo_rows(iso3: str, limit: int) -> tuple[list[dict[str, Any]], str]:
    workbook_path = SOURCE_ROOT / "GEO" / "Sub_nat_gmd.xlsx"
    if not workbook_path.exists():
        raise FileNotFoundError(f"missing source workbook: {workbook_path.relative_to(ROOT)}")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for GEO workbook extraction") from exc

    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return [], str(workbook_path.relative_to(ROOT))
    headers = [_to_text(col) for col in header_row]
    index = {name: idx for idx, name in enumerate(headers)}
    required = [
        "code",
        "byvar",
        "sample",
        "surveyid",
        "geo_year",
        "geo_source",
        "geo_level",
        "geo_idvar",
        "geo_id",
        "geo_nvar",
        "geo_name",
        "geo_code",
    ]
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(f"missing expected columns in GEO workbook: {missing}")

    rows: list[dict[str, Any]] = []
    for source_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        code = _to_text(values[index["code"]]).upper()
        if code != iso3:
            continue

        byvar = _to_text(values[index["byvar"]]).lower()
        level = _infer_subnat_level(byvar)
        geo_id = _to_text(values[index["geo_id"]])
        geo_code = _to_text(values[index["geo_code"]])

        survey_label = _to_text(values[index["sample"]]) or _to_text(values[index["geo_name"]])
        survey_variable = byvar or "subnatid"
        gmd_subnatidsurvey = (
            geo_code if survey_variable in {"subnatidsurvey", "subnatid"} else ""
        )

        row = {
            "survey_labels": survey_label,
            "survey_variables": survey_variable,
            "gmd_subnatid1": geo_code if level == 1 else "",
            "gmd_subnatid2": geo_code if level == 2 else "",
            "gmd_subnatid3": geo_code if level == 3 else "",
            "gmd_subnatid4": geo_code if level == 4 else "",
            "is_rep_subnat1": level == 1,
            "is_rep_subnat2": level == 2,
            "is_rep_subnat3": level == 3,
            "is_rep_subnat4": level == 4,
            "representative_level": level,
            "gmd_subnatidsurvey": gmd_subnatidsurvey,
            "geo_year": _to_text(values[index["geo_year"]]),
            "geo_source": _to_text(values[index["geo_source"]]),
            "geo_level": _to_text(values[index["geo_level"]]),
            "geo_idvar": _to_text(values[index["geo_idvar"]]),
            "geo_id": geo_id,
            "geo_nvar": _to_text(values[index["geo_nvar"]]),
            "geo_name": _to_text(values[index["geo_name"]]),
            "source_row": source_row,
        }
        rows.append(row)

    rows = _dedupe_geo_rows(rows)
    if limit > 0:
        rows = rows[:limit]

    return rows, str(workbook_path.relative_to(ROOT))


def list_geo_iso3s() -> tuple[list[str], str]:
    workbook_path = SOURCE_ROOT / "GEO" / "Sub_nat_gmd.xlsx"
    if not workbook_path.exists():
        return [], str(workbook_path.relative_to(ROOT))

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for GEO workbook extraction") from exc

    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return [], str(workbook_path.relative_to(ROOT))

    headers = [_to_text(col) for col in header_row]
    index = {name: idx for idx, name in enumerate(headers)}
    if "code" not in index:
        raise ValueError("missing expected column in GEO workbook: code")

    seen: set[str] = set()
    for values in ws.iter_rows(min_row=2, values_only=True):
        code = _to_text(values[index["code"]]).upper()
        if _is_valid_iso3(code):
            seen.add(code)

    return sorted(seen), str(workbook_path.relative_to(ROOT))


def run_geo_example(iso3: str, limit: int, write: bool) -> dict[str, Any]:
    iso3 = (iso3 or "").upper()
    if not _is_valid_iso3(iso3):
        return {"ok": False, "error": f"invalid iso3 '{iso3}'"}

    rows, source = extract_geo_rows(iso3, limit)
    if not rows:
        return {
            "ok": False,
            "error": f"no GEO rows found for {iso3}",
            "source": source,
        }

    country_name = _country_name_for_iso3(iso3)
    records = _build_geo_parameter_records(rows, source)

    written = None
    if write:
        out_path = _emit_geo_parameter_draft(iso3, country_name, rows, source)
        written = str(out_path.relative_to(ROOT))

    return {
        "ok": True,
        "iso3": iso3,
        "country_name": country_name,
        "source": source,
        "rows_returned": len(rows),
        "parameters": records,
        "written": written,
    }


def inspect_inputs() -> dict[str, list[str]]:
    isced = _find_xlsx(SOURCE_ROOT / "ISCED")
    jmp = _find_xlsx(SOURCE_ROOT / "JMP")
    geo = _find_xlsx(SOURCE_ROOT / "GEO")
    return {
        "ISCED": [str(path.relative_to(ROOT)) for path in isced],
        "JMP": [str(path.relative_to(ROOT)) for path in jmp],
        "GEO": [str(path.relative_to(ROOT)) for path in geo],
    }


def load_parameter_registry() -> dict[str, ParameterDefinition]:
    registry: dict[str, ParameterDefinition] = {}

    for path in sorted((ROOT / "knowledge" / "parameters").glob("*.md")):
        data, _ = load_markdown(path)
        parameter = ParameterDefinition.model_validate(data)
        if parameter.parameter_id in registry:
            raise ValueError(f"duplicate parameter_id in registry: {parameter.parameter_id}")
        registry[parameter.parameter_id] = parameter

    for path in sorted(CONTRACT_ROOT.glob("*.yaml")):
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError(f"contract is not a mapping: {path}")
        parameter = ParameterDefinition.model_validate(data)
        if parameter.parameter_id in registry:
            raise ValueError(f"duplicate parameter_id in registry: {parameter.parameter_id}")
        registry[parameter.parameter_id] = parameter

    return registry


def _extract_education(path: Path) -> tuple[str, str, list[dict[str, Any]], int | None]:
    module = load_edu_module()
    result = module.extract(str(path), out=None)
    instruction = result["instruction"]
    rows = transform_education_rows(result)
    return instruction.iso3, instruction.country, rows, instruction.source.reference_year


def _extract_wash(path: Path) -> tuple[str, str, list[dict[str, Any]], list[dict[str, Any]], dict[str, Any], dict[str, Any]]:
    module = load_wash_module()
    result = module.extract(str(path), out=None)
    water_rows = transform_wash_rows(result, "water")
    sanitation_rows = transform_wash_rows(result, "sanitation")

    water_instruction = result["domains"]["water"]["instruction"]
    bench_water = to_benchmark_payload(result, "water")
    bench_san = to_benchmark_payload(result, "sanitation")
    return (
        water_instruction.iso3,
        water_instruction.country,
        water_rows,
        sanitation_rows,
        bench_water,
        bench_san,
    )


def run_extract(*, force: bool = False) -> dict[str, Any]:
    inputs = inspect_inputs()
    if not inputs["ISCED"] and not inputs["JMP"] and not inputs["GEO"]:
        return {
            "ok": False,
            "errors": ["no input workbooks found under extraction/10_source/country-parameters-inputs"],
            "inputs": inputs,
            "written": [],
        }

    written: list[str] = []
    skipped: list[dict[str, Any]] = []
    incremental_skipped: list[dict[str, Any]] = []
    cleaned = _cleanup_stale_root_level_yaml()
    state = _load_incremental_state()
    extractor_sig = _extractor_signature()

    for relative in inputs["ISCED"]:
        source = ROOT / relative
        iso3, country_name, rows, _reference_year = _extract_education(source)
        if not _is_valid_iso3(iso3):
            skipped.append(
                {
                    "file": str(source.relative_to(ROOT)),
                    "parameter_id": "PARAM-EDU-LEVEL-CROSSWALK",
                    "reason": f"invalid iso3 '{iso3}'",
                }
            )
            continue
        if not rows:
            stale_path = DRAFT_ROOT / iso3 / "PARAM-EDU-LEVEL-CROSSWALK.yaml"
            _drop_artifact_state(state, f"draft:{iso3}:PARAM-EDU-LEVEL-CROSSWALK")
            if _delete_if_exists(stale_path):
                cleaned.append(str(stale_path.relative_to(ROOT)))
            skipped.append(
                {
                    "file": str(source.relative_to(ROOT)),
                    "parameter_id": "PARAM-EDU-LEVEL-CROSSWALK",
                    "reason": "no rows extracted",
                }
            )
            continue
        edu_out = DRAFT_ROOT / iso3 / "PARAM-EDU-LEVEL-CROSSWALK.yaml"
        edu_sig = _signature(
            _file_sha256(source),
            _parameter_schema_signature("PARAM-EDU-LEVEL-CROSSWALK"),
            extractor_sig,
        )
        edu_key = f"draft:{iso3}:PARAM-EDU-LEVEL-CROSSWALK"
        if not _should_emit(state=state, key=edu_key, signature=edu_sig, output_path=edu_out, force=force):
            incremental_skipped.append(
                {
                    "file": str(edu_out.relative_to(ROOT)),
                    "parameter_id": "PARAM-EDU-LEVEL-CROSSWALK",
                    "reason": "unchanged input/schema",
                }
            )
            continue
        out_path = emit_parameter_draft(
            iso3=iso3,
            country_name=country_name,
            parameter_id="PARAM-EDU-LEVEL-CROSSWALK",
            rows=rows,
            source=str(source.relative_to(ROOT)),
            effective_from=None,
            effective_to=None,
        )
        _record_artifact_state(state, edu_key, edu_sig, out_path)
        written.append(str(out_path.relative_to(ROOT)))

    for relative in inputs["JMP"]:
        source = ROOT / relative
        (
            iso3,
            country_name,
            water_rows,
            sanitation_rows,
            bench_water,
            bench_san,
        ) = _extract_wash(source)

        if not _is_valid_iso3(iso3):
            skipped.append(
                {
                    "file": str(source.relative_to(ROOT)),
                    "parameter_id": "PARAM-WASH-*",
                    "reason": f"invalid iso3 '{iso3}'",
                }
            )
            continue

        if water_rows:
            water_out = DRAFT_ROOT / iso3 / "PARAM-WASH-WATER-CROSSWALK.yaml"
            water_sig = _signature(
                _file_sha256(source),
                _parameter_schema_signature("PARAM-WASH-WATER-CROSSWALK"),
                extractor_sig,
            )
            water_key = f"draft:{iso3}:PARAM-WASH-WATER-CROSSWALK"
            if _should_emit(
                state=state,
                key=water_key,
                signature=water_sig,
                output_path=water_out,
                force=force,
            ):
                water_path = emit_parameter_draft(
                    iso3=iso3,
                    country_name=country_name,
                    parameter_id="PARAM-WASH-WATER-CROSSWALK",
                    rows=water_rows,
                    source=str(source.relative_to(ROOT)),
                    effective_from=None,
                    effective_to=None,
                )
                _record_artifact_state(state, water_key, water_sig, water_path)
                written.append(str(water_path.relative_to(ROOT)))
            else:
                incremental_skipped.append(
                    {
                        "file": str(water_out.relative_to(ROOT)),
                        "parameter_id": "PARAM-WASH-WATER-CROSSWALK",
                        "reason": "unchanged input/schema",
                    }
                )
        else:
            _drop_artifact_state(state, f"draft:{iso3}:PARAM-WASH-WATER-CROSSWALK")
            stale_water = DRAFT_ROOT / iso3 / "PARAM-WASH-WATER-CROSSWALK.yaml"
            if _delete_if_exists(stale_water):
                cleaned.append(str(stale_water.relative_to(ROOT)))
            skipped.append(
                {
                    "file": str(source.relative_to(ROOT)),
                    "parameter_id": "PARAM-WASH-WATER-CROSSWALK",
                    "reason": "no rows extracted",
                }
            )

        if sanitation_rows:
            san_out = DRAFT_ROOT / iso3 / "PARAM-WASH-SANITATION-CROSSWALK.yaml"
            san_sig = _signature(
                _file_sha256(source),
                _parameter_schema_signature("PARAM-WASH-SANITATION-CROSSWALK"),
                extractor_sig,
            )
            san_key = f"draft:{iso3}:PARAM-WASH-SANITATION-CROSSWALK"
            if _should_emit(
                state=state,
                key=san_key,
                signature=san_sig,
                output_path=san_out,
                force=force,
            ):
                san_path = emit_parameter_draft(
                    iso3=iso3,
                    country_name=country_name,
                    parameter_id="PARAM-WASH-SANITATION-CROSSWALK",
                    rows=sanitation_rows,
                    source=str(source.relative_to(ROOT)),
                    effective_from=None,
                    effective_to=None,
                )
                _record_artifact_state(state, san_key, san_sig, san_path)
                written.append(str(san_path.relative_to(ROOT)))
            else:
                incremental_skipped.append(
                    {
                        "file": str(san_out.relative_to(ROOT)),
                        "parameter_id": "PARAM-WASH-SANITATION-CROSSWALK",
                        "reason": "unchanged input/schema",
                    }
                )
        else:
            _drop_artifact_state(state, f"draft:{iso3}:PARAM-WASH-SANITATION-CROSSWALK")
            stale_san = DRAFT_ROOT / iso3 / "PARAM-WASH-SANITATION-CROSSWALK.yaml"
            if _delete_if_exists(stale_san):
                cleaned.append(str(stale_san.relative_to(ROOT)))
            skipped.append(
                {
                    "file": str(source.relative_to(ROOT)),
                    "parameter_id": "PARAM-WASH-SANITATION-CROSSWALK",
                    "reason": "no rows extracted",
                }
            )

        bench_water_key = f"benchmark:{iso3}:water"
        bench_water_path = ROOT / "governance" / "benchmarks" / f"{iso3}_water_benchmark_draft.yaml"
        bench_water_sig = _signature(_file_sha256(source), "water", extractor_sig)
        if _should_emit(
            state=state,
            key=bench_water_key,
            signature=bench_water_sig,
            output_path=bench_water_path,
            force=force,
        ):
            b1 = emit_benchmark(
                iso3,
                "water",
                bench_water,
                str(source.relative_to(ROOT)),
            )
            _record_artifact_state(state, bench_water_key, bench_water_sig, b1)
            written.append(str(b1.relative_to(ROOT)))
        else:
            incremental_skipped.append(
                {
                    "file": str(bench_water_path.relative_to(ROOT)),
                    "parameter_id": "BENCHMARK-WATER",
                    "reason": "unchanged input/schema",
                }
            )

        bench_san_key = f"benchmark:{iso3}:sanitation"
        bench_san_path = ROOT / "governance" / "benchmarks" / f"{iso3}_sanitation_benchmark_draft.yaml"
        bench_san_sig = _signature(_file_sha256(source), "sanitation", extractor_sig)
        if _should_emit(
            state=state,
            key=bench_san_key,
            signature=bench_san_sig,
            output_path=bench_san_path,
            force=force,
        ):
            b2 = emit_benchmark(
                iso3,
                "sanitation",
                bench_san,
                str(source.relative_to(ROOT)),
            )
            _record_artifact_state(state, bench_san_key, bench_san_sig, b2)
            written.append(str(b2.relative_to(ROOT)))
        else:
            incremental_skipped.append(
                {
                    "file": str(bench_san_path.relative_to(ROOT)),
                    "parameter_id": "BENCHMARK-SANITATION",
                    "reason": "unchanged input/schema",
                }
            )

    if inputs["GEO"]:
        try:
            geo_iso3s, geo_source = list_geo_iso3s()
        except Exception as exc:  # noqa: BLE001
            skipped.append(
                {
                    "file": "extraction/10_source/country-parameters-inputs/GEO/Sub_nat_gmd.xlsx",
                    "parameter_id": "PARAM-GEO-GMD-CROSSWALK",
                    "reason": f"failed to inspect GEO workbook: {exc}",
                }
            )
            geo_iso3s = []
            geo_source = "extraction/10_source/country-parameters-inputs/GEO/Sub_nat_gmd.xlsx"

        for iso3 in geo_iso3s:
            try:
                geo_rows, _ = extract_geo_rows(iso3, 0)
            except Exception as exc:  # noqa: BLE001
                skipped.append(
                    {
                        "file": geo_source,
                        "parameter_id": "PARAM-GEO-GMD-CROSSWALK",
                        "reason": f"{iso3}: extraction failed ({exc})",
                    }
                )
                continue

            if not geo_rows:
                stale_geo = DRAFT_ROOT / iso3 / "PARAM-GEO-GMD-CROSSWALK.yaml"
                _drop_artifact_state(state, f"draft:{iso3}:PARAM-GEO-GMD-CROSSWALK")
                if _delete_if_exists(stale_geo):
                    cleaned.append(str(stale_geo.relative_to(ROOT)))
                skipped.append(
                    {
                        "file": geo_source,
                        "parameter_id": "PARAM-GEO-GMD-CROSSWALK",
                        "reason": f"{iso3}: no rows extracted",
                    }
                )
                continue

            geo_out = DRAFT_ROOT / iso3 / "PARAM-GEO-GMD-CROSSWALK.yaml"
            geo_sig = _signature(
                _file_sha256(ROOT / geo_source),
                _parameter_schema_signature("PARAM-GEO-GMD-CROSSWALK"),
                extractor_sig,
                iso3,
            )
            geo_key = f"draft:{iso3}:PARAM-GEO-GMD-CROSSWALK"
            if not _should_emit(state=state, key=geo_key, signature=geo_sig, output_path=geo_out, force=force):
                incremental_skipped.append(
                    {
                        "file": str(geo_out.relative_to(ROOT)),
                        "parameter_id": "PARAM-GEO-GMD-CROSSWALK",
                        "reason": "unchanged input/schema",
                    }
                )
                continue

            geo_path = _emit_geo_parameter_draft(
                iso3,
                _country_name_for_iso3(iso3),
                geo_rows,
                geo_source,
            )
            _record_artifact_state(state, geo_key, geo_sig, geo_path)
            written.append(str(geo_path.relative_to(ROOT)))

    _save_incremental_state(state)

    return {
        "ok": True,
        "inputs": inputs,
        "written": written,
        "cleaned": cleaned,
        "skipped": skipped,
        "incremental_skipped": incremental_skipped,
        "incremental_state": str(STATE_PATH.relative_to(ROOT)),
        "force": force,
    }


def run_check() -> dict[str, Any]:
    base = DRAFT_ROOT
    errors: list[dict[str, str]] = []
    warnings: list[dict[str, str]] = []
    try:
        registry = load_parameter_registry()
    except Exception as exc:  # noqa: BLE001
        return {
            "ok": False,
            "count": 0,
            "files": [],
            "errors": [{"file": "registry", "error": str(exc)}],
            "warnings": [],
        }

    findings: list[dict[str, Any]] = []
    paths = [path for path in sorted(base.rglob("*.yaml")) if CONTRACT_ROOT not in path.parents]
    if not paths:
        return {
            "ok": False,
            "count": 0,
            "files": [],
            "errors": [
                {
                    "file": str(base.relative_to(ROOT)),
                    "error": "no draft country-parameter files found",
                }
            ],
            "warnings": [],
        }

    for path in paths:
        relative_path = path.relative_to(base)
        # Only validate drafts placed under an ISO3 folder (e.g., VNM/file.yaml).
        if len(relative_path.parts) < 2:
            warnings.append(
                {
                    "file": str(path.relative_to(ROOT)),
                    "warning": "skipped non-ISO3 root-level draft file",
                }
            )
            continue
        iso3_dir = relative_path.parts[0]
        if not _is_valid_iso3(iso3_dir):
            warnings.append(
                {
                    "file": str(path.relative_to(ROOT)),
                    "warning": f"skipped draft under non-ISO3 folder '{iso3_dir}'",
                }
            )
            continue

        try:
            text = path.read_text(encoding="utf-8")
        except FileNotFoundError:
            warnings.append(
                {
                    "file": str(path.relative_to(ROOT)),
                    "warning": "file disappeared during scan; skipped",
                }
            )
            continue
        file_ok = True
        detail_error = ""
        try:
            payload = yaml.safe_load(text)
            if not isinstance(payload, dict):
                raise ValueError("draft file must contain a mapping")
            parsed = CountryParameterFile.model_validate(payload, context={"registry": registry})

            for record in parsed.parameters:
                definition = registry[record.parameter_id]
                if definition.value_type == "table" and len(record.value) == 0:
                    warnings.append(
                        {
                            "file": str(path.relative_to(ROOT)),
                            "warning": f"{record.parameter_id} has empty table value; skipped as non-blocking draft",
                        }
                    )
        except Exception as exc:  # noqa: BLE001
            file_ok = False
            detail_error = str(exc)
            errors.append(
                {
                    "file": str(path.relative_to(ROOT)),
                    "error": detail_error,
                }
            )

        findings.append(
            {
                "file": str(path.relative_to(ROOT)),
                "size": len(text),
                "has_parameters": "parameters:" in text,
                "has_value": "value:" in text,
                "valid": file_ok,
                "error": detail_error,
            }
        )
    return {
        "ok": len(errors) == 0,
        "count": len(findings),
        "files": findings,
        "errors": errors,
        "warnings": warnings,
    }


def run_bulk(*, force: bool = False) -> dict[str, Any]:
    return run_extract(force=force)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Country parameter input extraction")
    parser.add_argument("command", choices=["inspect", "extract", "check", "bulk", "geo-example"])
    parser.add_argument("--iso3", default="VNM", help="ISO3 filter for geo-example command")
    parser.add_argument("--limit", type=int, default=12, help="Max rows to return for geo-example")
    parser.add_argument(
        "--write",
        action="store_true",
        help="Write geo-example rows as a draft YAML under extraction/20_drafts/runs/country-parameters/<ISO3>/",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Force rewrite all outputs for extract/bulk, bypassing incremental checks.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.command == "inspect":
        print(json.dumps(inspect_inputs(), indent=2, ensure_ascii=False))
        return 0
    if args.command == "extract":
        result = run_extract(force=args.force)
        print(json.dumps(result, indent=2, ensure_ascii=False))
        return 0 if result.get("ok", False) else 1
    if args.command == "check":
        result = run_check()
        print(json.dumps(result, indent=2, ensure_ascii=False))
        return 0 if result.get("ok", False) else 1
    if args.command == "bulk":
        result = run_bulk(force=args.force)
        print(json.dumps(result, indent=2, ensure_ascii=False))
        return 0 if result.get("ok", False) else 1
    if args.command == "geo-example":
        result = run_geo_example(args.iso3, args.limit, args.write)
        print(json.dumps(result, indent=2, ensure_ascii=False))
        return 0 if result.get("ok", False) else 1
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
